"""Auditoría de solo lectura del corpus; escribe exclusivamente junto a este script.
Uso desde cualquier directorio: python ruta/auditar.py
No modifica maestros, fuentes, documentos ni scripts de producción.
"""
from pathlib import Path
from collections import Counter
import ast
import hashlib
import importlib.util
import json
import re
import sys
import wave
import pandas as pd
from docx import Document

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent.parent
sys.stdout.reconfigure(encoding="utf-8")
sys.dont_write_bytecode = True
sys.path.insert(0, str(ROOT / "scripts"))

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def load(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / (name + ".py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

def main():
    files = sorted(p for p in ROOT.rglob("*") if p.is_file()
                   and ".git" not in p.parts and "__pycache__" not in p.parts
                   and OUT not in p.parents)
    inv = pd.DataFrame([{"archivo": p.relative_to(ROOT).as_posix(), "bytes": p.stat().st_size,
                         "sha256": sha(p)} for p in files])
    sheets, docs, python_checks = [], [], []
    for p in files:
        rel = p.relative_to(ROOT).as_posix()
        if p.suffix == ".xlsx":
            try:
                x = pd.ExcelFile(p)
                for s in x.sheet_names:
                    d = x.parse(s, dtype=str).fillna("")
                    sheets.append({"archivo": rel, "hoja": s, "filas": len(d),
                                   "columnas": " | ".join(map(str, d.columns))})
                x.close()
            except Exception as e:
                sheets.append({"archivo": rel, "error": repr(e)})
        if p.suffix == ".py":
            try:
                ast.parse(p.read_text(encoding="utf-8-sig"))
                python_checks.append({"archivo": rel, "sintaxis": "OK"})
            except Exception as e:
                python_checks.append({"archivo": rel, "sintaxis": repr(e)})
        if p.suffix == ".docx":
            doc = Document(p)
            text = "\n".join([x.text for x in doc.paragraphs] +
                             [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
            ids = re.findall(r"\b(?:GEN|CLI)-\d{4}\b", text)
            docs.append({"archivo": rel, "tablas": len(doc.tables), "IDs_unicos": len(set(ids)),
                         "Rimay_n": len(re.findall(r"Rimay \d+", text)),
                         "FALTA_VALIDAR": text.count("FALTA VALIDAR"), "sha256": sha(p)})
            dst = OUT / "texto_documentos" / (rel.replace("/", "__") + ".txt")
            dst.parent.mkdir(exist_ok=True)
            dst.write_text(text, encoding="utf-8")
    tables = pd.read_excel(ROOT / "02_maestro/corpus_maestro.xlsx", sheet_name=None, dtype=str)
    tables = {k: v.fillna("") for k, v in tables.items()}
    c, o, qa = tables["Corpus"], tables["Origen"], tables["QA"]
    merged = c.merge(o, on="ID", validate="one_to_one")
    selected = merged[merged.Bloque.ne("")].copy()
    block_rows, differences = [], []
    for b, g in selected.groupby("Bloque"):
        block_rows.append({"Bloque": b, "filas": len(g), "General": g.Dominio.eq("General").sum(),
            "Médico": g.Dominio.eq("Médico").sum(), "sin_borrador": g.Quechua_borrador.eq("").sum(),
            "Por_revisar": g.Estado.eq("Por_revisar").sum(),
            "validados": g.Quechua_validado.ne("").sum(),
            "Test_sugerido": g.Test_sugerido.eq("Sí").sum()})
        info = g.set_index("ID")
        doc = Document(ROOT / f"03_bloques/{b}/{b}_traductor.docx")
        found = []
        for t in doc.tables:
            for r in t.rows:
                v = [x.text.strip() for x in r.cells]
                if len(v) != 6 or not re.fullmatch(r"(?:GEN|CLI)-\d{4}", v[1]):
                    continue
                i = v[1]; found.append(i)
                if i not in info.index:
                    differences.append({"Bloque": b, "ID": i, "problema": "ID ajeno"}); continue
                row = info.loc[i]
                expected = "" if row.Control == "Sí" else row.Quechua_borrador
                for field, got, want in [("Español", v[3], row.Español), ("Quechua", v[4], expected)]:
                    if got != want:
                        differences.append({"Bloque": b, "ID": i, "problema": field, "Word": got, "Maestro": want})
        for i in set(info.index) - set(found):
            differences.append({"Bloque": b, "ID": i, "problema": "falta en Word"})
        for i, n in Counter(found).items():
            if n > 1: differences.append({"Bloque": b, "ID": i, "problema": f"duplicado {n}"})
    # Prueba real de QA de audio: PCM silencioso, sin usar ningún audio de producción.
    wav = OUT / "PRUEBA_SILENCIO_NO_USAR.wav"
    with wave.open(str(wav), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(16000)
        w.writeframes(b"\x00\x00" * 32000)
    duration, issues = load("08_manifest").revisar_wav(wav, "Buenos días")
    from comun import clave_normalizada
    keys = [{"a": a, "b": b, "colision": clave_normalizada(a) == clave_normalizada(b)}
            for a, b in [("¿Tiene fiebre?", "Tiene fiebre."), ("El médico", "El medicó"), ("Sí", "Si")]]
    # Lista explícita de observaciones editoriales; no son correcciones automáticas.
    ids_review = ["GEN-0011", "GEN-0533", "GEN-0030", "GEN-0056", "CLI-0364", "CLI-0287", "CLI-0624"]
    review = merged[merged.ID.isin(ids_review)].copy()
    review["Observacion_revision"] = review.ID.map({
        "GEN-0011": "Frase literaria larga con varias subordinadas; revisar facilidad de grabación y contexto.",
        "GEN-0533": "Mi lengua vive en mi cabeza: posible metáfora o frase artificial; aclarar intención antes de traducir.",
        "GEN-0030": "Referencia a «mi cordero»: decidir si sirve al dominio general previsto; no es un error por sí misma.",
        "GEN-0056": "Hay que esperar: validar valor modal y sentido; QA existente advierte repetición en borrador.",
        "CLI-0287": "QA existente advierte discrepancia después del nacimiento / al enfermarse; resolver con traductor, conservando contexto pediátrico.",
        "CLI-0624": "Revisar par con CLI-1140: el español del médico dice líquidos y la respuesta menciona mate; evaluación clínica y de alineación antes de usarlo como modelo.",
        "CLI-0364": "¿Te sacude el cuerpo?: revisar naturalidad y significado clínico con hablante y profesional sanitario."})
    summary = {"archivos_inventariados": len(files), "hojas_excel": len(sheets), "documentos_word": len(docs),
        "corpus": len(c), "origen": len(o), "ids_duplicados": int(c.ID.duplicated().sum()),
        "estados": c.Estado.value_counts().to_dict(), "splits": c.Split.value_counts().to_dict(),
        "quechua_validado_no_vacio": int(c.Quechua_validado.ne("").sum()),
        "seleccionados": len(selected), "sin_bloque": int(c.Bloque.eq("").sum()),
        "diferencias_word_maestro": len(differences),
        "prueba_silencio": {"duracion": duration, "avisos": issues, "limitacion_confirmada": not issues},
        "colisiones_clave": keys}
    (OUT / "resumen.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    source_summary = selected.groupby(["Bloque", "Fuente", "Tipo_borrador"]).size().reset_index(name="n")
    with pd.ExcelWriter(OUT / "AUDITORIA_CORPUS.xlsx", engine="openpyxl") as writer:
        for name, df in [("Bloques", pd.DataFrame(block_rows)), ("Fuentes_por_bloque", source_summary),
                         ("B02_revision", merged[merged.Bloque.eq("B02")].merge(qa, on="ID", how="left")),
                         ("Observaciones_editoriales", review), ("Documentos", pd.DataFrame(docs)),
                         ("Diferencias_Word", pd.DataFrame(differences)), ("Hojas_Excel", pd.DataFrame(sheets)),
                         ("Inventario_SHA256", inv), ("Sintaxis_python", pd.DataFrame(python_checks)),
                         ("Colisiones_clave", pd.DataFrame(keys))]:
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]; ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    # Comprobar que ningún archivo previamente inventariado cambió al ejecutar la auditoría.
    changed = [r.archivo for r in inv.itertuples() if sha(ROOT / r.archivo) != r.sha256]
    (OUT / "integridad.json").write_text(json.dumps({"archivos_originales_modificados": changed}, indent=2), encoding="utf-8")
    assert not changed, changed
    print(json.dumps(summary, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
