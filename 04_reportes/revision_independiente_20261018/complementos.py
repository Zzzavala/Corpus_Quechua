"""Diagnósticos adicionales sin modificar entradas: versiones, paralelos y gráficos."""
from pathlib import Path
import hashlib
import json
import re
import pandas as pd
import pypdfium2 as pdfium
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent.parent

def main():
    original = json.loads((ROOT / "04_reportes/huella_fuentes_F0.json").read_text(encoding="utf-8"))
    hashes = []
    for name, v in original.items():
        data = (ROOT / name).read_bytes()
        hashes.append({"archivo": name, "igual_bytes": hashlib.sha256(data).hexdigest() == v["sha256"],
                       "igual_normalizando_CRLF_a_LF": hashlib.sha256(data.replace(b"\r\n", b"\n")).hexdigest() == v["sha256"]})
    pairs = []
    for pe in (ROOT / "Fuentes/AmericasNLP").rglob("*.es"):
        pq = pe.with_suffix(".quy")
        e = pe.read_text(encoding="utf-8-sig").splitlines()
        q = pq.read_text(encoding="utf-8-sig").splitlines() if pq.exists() else []
        pairs.append({"es": pe.relative_to(ROOT).as_posix(), "n_es": len(e), "n_qu": len(q),
                      "vacias_es": sum(not x.strip() for x in e), "vacias_qu": sum(not x.strip() for x in q),
                      "mismo_numero_lineas": len(e)==len(q),
                      "nota": "mismo número no demuestra alineación semántica" if pq.exists() else "solo español; no inventar pareja"})
    pdfs = []
    for p in sorted((ROOT / "Fuentes").glob("*.pdf")):
        pdf = pdfium.PdfDocument(p); counts=[]
        for i in range(len(pdf)):
            page=pdf[i];tp=page.get_textpage();counts.append(len(tp.get_text_range().strip()));tp.close();page.close()
        pdfs.append({"archivo": p.name, "paginas": len(pdf), "paginas_sin_texto_extraible": sum(n==0 for n in counts)})
        pdf.close()
    versions = []
    for b in [f"B{i:02d}" for i in range(1,6)]:
        doc = Document(ROOT / f"03_bloques/{b}/{b}_traductor.docx")
        current=[]
        for t in doc.tables:
            for r in t.rows:
                c=[x.text for x in r.cells]
                if len(c)==6 and re.fullmatch(r"(?:GEN|CLI)-\d{4}",c[1]): current.append((c[1],c[3]))
        pdf=pdfium.PdfDocument(ROOT / f"04_reportes/revision_visual/{b}_traductor.pdf")
        text=""
        for i in range(len(pdf)):
            page=pdf[i];tp=page.get_textpage();text += tp.get_text_range()+"\n";tp.close();page.close()
        pdf.close()
        ids_pdf=re.findall(r"\b(?:GEN|CLI)-\d{4}\b",text)
        ids_current=[x[0] for x in current]
        versions.append({"Bloque":b,"IDs_Word":len(ids_current),"IDs_PDF":len(ids_pdf),
            "mismos_IDs_en_orden":ids_pdf==ids_current,"solo_en_Word":sorted(set(ids_current)-set(ids_pdf)),
            "primer_ID_Word":ids_current[0],"primer_ID_PDF":ids_pdf[0] if ids_pdf else None})
    master=pd.read_excel(ROOT / "02_maestro/corpus_maestro.xlsx",sheet_name=None,dtype=str)
    c=master["Corpus"].fillna("");o=master["Origen"].fillna("");m=c.merge(o,on="ID",validate="one_to_one")
    selected=m[m.Bloque.ne("")].copy()
    state=pd.crosstab(selected.Bloque,selected.Estado)
    source=pd.crosstab(selected.Bloque,selected.Fuente)
    fig,axes=plt.subplots(2,1,figsize=(11,9))
    state.plot.bar(stacked=True,ax=axes[0],rot=0)
    axes[0].set_title("Estado de los 1 000 candidatos: aún no hay quechua validado")
    axes[0].set_ylabel("Oraciones");axes[0].legend(loc="upper left",bbox_to_anchor=(1,1))
    source.plot.bar(stacked=True,ax=axes[1],rot=0,colormap="tab20")
    axes[1].set_title("La proporción 70/30 no asegura la misma mezcla de fuentes")
    axes[1].set_ylabel("Oraciones");axes[1].legend(loc="upper left",bbox_to_anchor=(1,1),fontsize=8)
    fig.tight_layout();fig.savefig(OUT/"estado_y_fuentes.png",dpi=150);plt.close(fig)
    results={"huellas_historicas":hashes,"paralelos":pairs,"pdf_fuentes":pdfs,"versiones_render":versions,
        "versiones_raiz":{b: (ROOT/f"{b}_traductor.docx").read_bytes()==(ROOT/f"03_bloques/{b}/{b}_traductor.docx").read_bytes() for b in ["B01","B02"]}}
    (OUT/"complementos.json").write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(results,ensure_ascii=True,indent=2))

if __name__=="__main__": main()
