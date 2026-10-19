"""Casos adversariales reales, exclusivamente sobre copias dentro de esta carpeta.
Los documentos y textos producidos son SIMULADOS, nunca material de grabación.
"""
from pathlib import Path
import hashlib
import json
import os
import subprocess
import sys
import openpyxl
from docx import Document

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent.parent
TMP = OUT / "PRUEBAS_SIMULADAS_NO_GRABAR"
TMP.mkdir(exist_ok=True)
env = dict(os.environ, PYTHONPATH=str(OUT / "dependencias"), PYTHONDONTWRITEBYTECODE="1", PYTHONIOENCODING="utf-8")

def run(script, *args):
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / script), *map(str, args)],
                       cwd=ROOT, env=env, capture_output=True, encoding="utf-8")
    return {"codigo": r.returncode, "salida": r.stdout, "error": r.stderr}

def main():
    original = ROOT / "02_maestro/corpus_maestro.xlsx"
    before = hashlib.sha256(original.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(original)
    ws = wb["Corpus"]; col = {c.value: c.column for c in ws[1]}
    rows = [r for r in range(2, ws.max_row+1) if ws.cell(r, col["Bloque"]).value == "B01"]
    target = rows[0]
    # Se reduce el bloque de prueba a un ID, sin alterar producción.
    for r in rows[1:]: ws.cell(r, col["Bloque"]).value = ""
    cid = ws.cell(target, col["ID"]).value
    es = ws.cell(target, col["Español"]).value
    ws.cell(target, col["Estado"]).value = "Validado"
    ws.cell(target, col["Quechua_validado"]).value = "SIMULADO anterior; NO GRABAR"
    master = TMP / "maestro_simulado.xlsx"; wb.save(master)

    def word(name, entries):
        d = Document(); d.add_paragraph("PRUEBA SIMULADA — NO USAR PARA GRABACIÓN")
        t = d.add_table(rows=0, cols=6)
        for e, q in entries:
            cells = t.add_row().cells
            for j, v in enumerate(["1", cid, "", e, q, ""]): cells[j].text = v
        p = TMP / name; d.save(p); return p

    common = ["B01", "--maestro", master, "--trabajo", TMP / "no_existe.xlsx", "--reporte", TMP / "importacion.md"]
    changed_doc = word("cambio_es_simulado.docx", [(es + " CAMBIO PROPUESTO", "SIMULADO propuesta")])
    result = {"ID": cid, "cambio_es": run("06_importar_traductor.py", *common, "--docx", changed_doc)}
    check = openpyxl.load_workbook(master)["Corpus"]
    result["tras_cambio"] = {k: check.cell(target,col[k]).value for k in ["Estado","Quechua_validado"]}
    result["grabacion_con_por_revisar"] = run("07_hoja_grabacion.py", "B01", "--maestro", master,
        "--salida", TMP / "grabacion_SIMULADA_NO_USAR.docx")
    dupdoc = word("duplicados_simulados.docx", [(es, "SIMULADO primera opción"), (es, "SIMULADO segunda opción incompatible")])
    result["duplicados"] = run("06_importar_traductor.py", *common, "--docx", dupdoc)
    check = openpyxl.load_workbook(master)["Corpus"]
    result["tras_duplicados"] = {k: check.cell(target,col[k]).value for k in ["Estado","Quechua_validado"]}
    result["maestro_original_inalterado"] = hashlib.sha256(original.read_bytes()).hexdigest() == before
    result["casos_reproducidos"] = {
        "estado_por_revisar_pasa_a_hoja": result["cambio_es"]["codigo"] == 0 and result["tras_cambio"]["Estado"] == "Por_revisar" and result["grabacion_con_por_revisar"]["codigo"] == 0,
        "duplicado_conflictivo_valida_primero": result["duplicados"]["codigo"] == 0 and result["tras_duplicados"]["Estado"] == "Validado" and result["tras_duplicados"]["Quechua_validado"] == "SIMULADO primera opción"}
    (OUT / "pruebas_importacion.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    assert result["maestro_original_inalterado"]

if __name__ == "__main__": main()
