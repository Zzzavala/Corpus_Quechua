"""Ida y vuelta Word → Excel (§8.5): importa el Word devuelto por el traductor.

1. Empareja POR ID (no por posición) y reporta IDs faltantes, sobrantes o duplicados.
2. Escribe el quechua final en el maestro (Quechua_validado, Estado = Validado) y en BXX_trabajo.xlsx, y deja cada cambio
   en la pestaña Cambios (antes, después, regla, fecha): reimportar nunca sobrescribe sin rastro.
3. Separa el comentario del traductor del «Aviso:» que venía prellenado.
4. Métricas: sin cambios / editado / reescrito y distancia de edición; análisis de las oraciones de control (D7).
   Se guardan en la pestaña Traductor_metricas del maestro.
5. Revisa la ortografía del traductor con las reglas de §10.3 SOLO para avisar (nunca corrige).

Protecciones:
- Si el Word tiene «Control de cambios» sin aceptar, no se importa (el texto insertado o borrado se leería mal).
- Si el traductor cambió el español de una fila, esa fila NO se valida: queda Revisar_Juan para que Juan decida.
- Si un sublote vuelve 100 % igual al borrador (con al menos traductor.min_filas_sublote_sin_tocar filas), no se valida
  solo: puede no haberse revisado. Se confirma con el traductor y se reimporta con --confirmar-sublotes B01-S2.

Uso:
    python scripts/06_importar_traductor.py B01 --marcar-enviado
        (al mandar el Word: marca las filas del bloque como Enviado; desde ahí f6 ya no rearma ese bloque)
    python scripts/06_importar_traductor.py B01
        (lee 03_bloques/B01/B01_traductor_devuelto.docx)
    python scripts/06_importar_traductor.py B01 --confirmar-sublotes B01-S4
    python scripts/06_importar_traductor.py B01 --docx RUTA --maestro RUTA --trabajo RUTA --reporte RUTA
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from rapidfuzz.distance import Levenshtein

from comun import CFG, RAIZ, escribir_texto, ruta, texto
import planilla as P
from qa import aviso_word
import quechua as Q

RE_ID = re.compile(r"^(GEN|CLI)-\d{4}$")
CT = CFG["traductor"]
REGLA = "importación del Word del traductor (06)"


def norm_cmp(s: str) -> str:
    s = texto(s).lower().replace("’", "'")
    s = re.sub(r"[^\w\s']", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_es(s: str) -> str:
    return re.sub(r"\s+", " ", texto(s)).strip()


def distancia(a: str, b: str) -> float:
    return Levenshtein.normalized_distance(norm_cmp(a), norm_cmp(b))


def control_de_cambios(doc) -> tuple[int, int]:
    """Cantidad de inserciones/borrados pendientes (Control de cambios) y de comentarios de Word."""
    cuerpo = doc.element.body
    return len(cuerpo.xpath(".//w:ins | .//w:del")), len(cuerpo.xpath(".//w:commentReference"))


def leer_word(doc) -> list[dict]:
    filas = []
    for t in doc.tables:
        for row in t.rows:
            celdas = row.cells
            if len(celdas) < 6:
                continue
            idv = celdas[1].text.strip()
            if not RE_ID.match(idv):
                continue
            filas.append({"ID": idv, "Español": celdas[3].text.strip(), "Quechua": celdas[4].text.strip(),
                          "Observaciones": celdas[5].text.strip()})
    return filas


def separar_comentario(obs: str, aviso: str) -> str:
    obs = obs.strip()
    if aviso and obs.startswith(aviso):
        return obs[len(aviso):].lstrip(" .;:-—\n").rstrip()
    if obs.startswith("Aviso:"):
        partes = obs.split("\n", 1)
        if len(partes) == 2:
            return partes[1].strip()
        m = re.search(r"(Comentario|Obs\.?|Nota)\s*:", obs, re.I)
        return obs[m.start():].strip() if m else ""
    return obs


def texto_anclaje(d_med, d_ctrl, n_ctrl: int, n_cb: int) -> str:
    """D7: compara cuánto cambia el traductor un borrador visible (d_med) con cuánto difiere su traducción libre del
    borrador que no vio (d_ctrl). Si casi no cambia lo visible pero lo libre sale muy distinto, hay anclaje."""
    if not n_ctrl or not n_cb or d_ctrl == "" or d_med == "":
        return "Sin oraciones de control en este bloque."
    if d_ctrl == 0:
        return ("Las traducciones libres de control coinciden con el borrador oculto (distancia 0): los borradores parecen "
                "buenos; no se puede medir anclaje con este bloque.")
    razon = d_med / d_ctrl
    if razon < 0.5:
        return (f"Posible anclaje: con borrador cambia {d_med:.2f} y su traducción libre difiere {d_ctrl:.2f} del borrador "
                f"(razón {razon:.2f}).")
    return f"Sin señal clara de anclaje (con borrador {d_med:.2f} vs. libre {d_ctrl:.2f}; razón {razon:.2f})."


def agregar_cambio(ws_cam, cid: str, codigo: str, campo: str, antes: str, despues: str, regla: str, fecha: str):
    cols = P.columnas(ws_cam)
    fila = ws_cam.max_row + 1
    for k, v in {"ID": cid, "Código_original": codigo, "Campo": campo, "Antes": antes, "Después": despues, "Regla": regla,
                 "Fecha": fecha}.items():
        if k in cols:
            ws_cam.cell(row=fila, column=cols[k]).value = v


def marcar_enviado(b: str, p_maestro: Path, p_trabajo: Path) -> int:
    wb_m = P.abrir(p_maestro)
    ws_c = wb_m["Corpus"]
    corpus = {r["ID"]: r for r in P.leer(ws_c)}
    codigos = {r["ID"]: r.get("Código_original", "") for r in P.leer(wb_m["Origen"])}
    filas_c = P.filas_por_id(ws_c)
    fecha = datetime.now().strftime("%Y-%m-%d")
    ids = [i for i, r in corpus.items() if r["Bloque"] == b and r["Estado"] not in ("Validado", "Grabado", "Enviado")]
    if not ids:
        print(f"{b}: no hay filas para marcar (o ya están Enviado/Validado/Grabado).")
        return 0
    for i in ids:
        agregar_cambio(wb_m["Cambios"], i, codigos.get(i, ""), "Estado", corpus[i]["Estado"], "Enviado",
                       "Word enviado al traductor (06 --marcar-enviado)", fecha)
        P.poner(ws_c, filas_c[i], "Estado", "Enviado")
    P.guardar(wb_m, p_maestro)
    if p_trabajo.exists():
        wb_t = P.abrir(p_trabajo)
        ws_b = wb_t["Bloque"]
        fb = P.filas_por_id(ws_b)
        for i in ids:
            if i in fb:
                P.poner(ws_b, fb[i], "Estado", "Enviado")
        P.guardar(wb_t, p_trabajo)
    print(f"{b}: {len(ids)} filas marcadas como Enviado. Desde ahora f6 conserva la composición de {b} (congelado).")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("bloque")
    ap.add_argument("--docx")
    ap.add_argument("--maestro")
    ap.add_argument("--trabajo")
    ap.add_argument("--reporte")
    ap.add_argument("--marcar-enviado", action="store_true", help="marcar las filas del bloque como Enviado y salir")
    ap.add_argument("--confirmar-sublotes", default="",
                    help="sublotes que volvieron sin cambios y el traductor confirmó que sí revisó (p. ej. B01-S2,B01-S4)")
    a = ap.parse_args()
    b = a.bloque.upper()
    dir_b = ruta("bloques") / b

    def ruta_arg(valor, defecto: Path) -> Path:
        if not valor:
            return defecto
        p = Path(valor)
        return p if p.is_absolute() else RAIZ / p

    p_docx = ruta_arg(a.docx, dir_b / f"{b}_traductor_devuelto.docx")
    p_maestro = ruta_arg(a.maestro, RAIZ / CFG["rutas"]["maestro"])
    p_trabajo = ruta_arg(a.trabajo, dir_b / f"{b}_trabajo.xlsx")
    p_reporte = ruta_arg(a.reporte, ruta("reportes") / f"importacion_{b}.md")
    if a.marcar_enviado:
        return marcar_enviado(b, p_maestro, p_trabajo)
    confirmados = {s.strip().upper() for s in a.confirmar_sublotes.split(",") if s.strip()}
    if not p_docx.exists():
        print(f"No existe {p_docx}. Guarda ahí el Word que devolvió el traductor.")
        return 2

    doc = Document(p_docx)
    n_cc, n_com = control_de_cambios(doc)
    if n_cc:
        print(f"El Word tiene {n_cc} cambios sin aceptar (Control de cambios). No se importa: ábrelo en Word, "
              "Revisar → Aceptar todos los cambios, guarda y vuelve a correr.")
        return 4
    devuelto = leer_word(doc)
    wb_m = P.abrir(p_maestro)
    ws_c = wb_m["Corpus"]
    corpus = {r["ID"]: r for r in P.leer(ws_c)}
    codigos = {r["ID"]: r.get("Código_original", "") for r in P.leer(wb_m["Origen"])}
    avisos = {}
    for r in P.leer(wb_m["QA"]):
        avisos.setdefault(r["ID"], []).append((r["Tipo"], r["Severidad"], r["Detalle"]))
    esperados = [i for i, r in corpus.items() if r["Bloque"] == b]
    esperados.sort(key=lambda i: int(corpus[i]["N_bloque"] or 0))

    cuenta = Counter(f["ID"] for f in devuelto)
    duplicados = sorted(i for i, n in cuenta.items() if n > 1)
    faltantes = [i for i in esperados if i not in cuenta]
    sobrantes = sorted(i for i in cuenta if i not in set(esperados))
    por_id = {}
    for f in devuelto:
        por_id.setdefault(f["ID"], f)            # si hay duplicado, vale la primera aparición (se reporta)

    resultados = []
    for i in esperados:
        if i not in por_id:
            continue
        r = corpus[i]
        f = por_id[i]
        borrador = r["Quechua_borrador"]
        control = r.get("Control", "No") == "Sí"
        aviso = "" if (control or not borrador) else aviso_word(avisos.get(i, []))
        comentario = separar_comentario(f["Observaciones"], aviso)
        final = f["Quechua"]
        d = distancia(borrador, final) if (borrador and final) else None
        if not final:
            clase = "vacío"
        elif control:
            clase = "control"
        elif not borrador:
            clase = "sin borrador"
        elif d <= CT["umbral_sin_cambios"]:
            clase = "sin cambios"
        elif d <= CT["umbral_reescrito"]:
            clase = "editado"
        else:
            clase = "reescrito"
        orto = [det for _, det in Q.detectar_orto(final)] + Q.detectar_variedad(final) if final else []
        resultados.append({"ID": i, "final": final, "clase": clase, "d": d, "control": control, "comentario": comentario,
                           "orto": orto, "español_nuevo": f["Español"] if norm_es(f["Español"]) != norm_es(r["Español"]) else "",
                           "sublote": r["Sublote"]})

    # ------------------------------------------------------------- sublotes devueltos sin tocar (H7)
    por_sub = defaultdict(list)
    for x in resultados:
        if x["clase"] in ("sin cambios", "editado", "reescrito"):
            por_sub[x["sublote"]].append(x["clase"])
    sin_tocar = sorted(s for s, cl in por_sub.items()
                       if len(cl) >= CT["min_filas_sublote_sin_tocar"] and all(c == "sin cambios" for c in cl) and s not in confirmados)

    # ------------------------------------------------------------- métricas
    con_borr = [x for x in resultados if x["clase"] in ("sin cambios", "editado", "reescrito")]
    ctrl = [x for x in resultados if x["clase"] == "control"]
    n_cb = len(con_borr)
    pct = lambda c: round(100 * sum(1 for x in con_borr if x["clase"] == c) / n_cb, 1) if n_cb else 0.0
    d_med = round(sum(x["d"] for x in con_borr) / n_cb, 3) if n_cb else ""
    d_ctrl = round(sum(x["d"] for x in ctrl if x["d"] is not None) / len(ctrl), 3) if ctrl else ""
    anclaje = texto_anclaje(d_med, d_ctrl, len(ctrl), n_cb)
    recomendacion = ("Dejar de dar borradores: reescribe más del 50 % (no ahorran tiempo)." if n_cb and pct("reescrito") > 100 * CT["umbral_dejar_borradores"]
                     else "Se pueden seguir dando borradores.")
    fila_tm = {"Bloque": b, "Fecha_importación": datetime.now().strftime("%Y-%m-%d %H:%M"), "Filas": len(resultados),
               "Filas_con_borrador": n_cb, "Sin_cambios_%": pct("sin cambios"), "Editado_%": pct("editado"), "Reescrito_%": pct("reescrito"),
               "Distancia_media_borrador": d_med, "Control_n": len(ctrl), "Distancia_media_control": d_ctrl, "Anclaje": anclaje,
               "Recomendación": recomendacion, "Faltantes": ", ".join(faltantes), "Sobrantes": ", ".join(sobrantes),
               "Duplicados": ", ".join(duplicados)}

    # ------------------------------------------------------------- decidir qué se valida
    for x in resultados:
        if not x["final"]:
            x["accion"] = "no (vacía)"
        elif x["español_nuevo"]:
            x["accion"] = "Revisar_Juan (el traductor cambió el español)"
        elif x["sublote"] in sin_tocar and x["clase"] == "sin cambios":
            x["accion"] = "no (sublote devuelto sin cambios: confirmar)"   # lo que el traductor escribió (control, vacías) sí vale
        else:
            x["accion"] = "Validado"

    # ------------------------------------------------------------- escribir maestro (+ Cambios)
    filas_c = P.filas_por_id(ws_c)
    ws_cam = wb_m["Cambios"]
    fecha = datetime.now().strftime("%Y-%m-%d")
    # al reimportar no se repite lo que ya quedó registrado (mismo ID, campo y valor)
    ya = {(c["ID"], c["Campo"], c["Después"]) for c in P.leer(ws_cam) if REGLA in c.get("Regla", "")}

    def registrar(i, campo, antes, despues, regla=REGLA):
        if (i, campo, despues) not in ya:
            agregar_cambio(ws_cam, i, codigos.get(i, ""), campo, antes, despues, regla, fecha)
            ya.add((i, campo, despues))

    validados = 0
    for x in resultados:
        i, fc = x["ID"], filas_c[x["ID"]]
        r = corpus[i]
        if x["comentario"]:
            registrar(i, "Comentario_traductor", "", x["comentario"])
        if x["accion"] == "Validado":
            if r["Quechua_validado"] != x["final"]:
                registrar(i, "Quechua_validado", r["Quechua_validado"] or r["Quechua_borrador"], x["final"],
                          REGLA + (" — reimportación: reemplaza un validado anterior" if r["Quechua_validado"] else ""))
            if r["Estado"] != "Validado":
                agregar_cambio(ws_cam, i, codigos.get(i, ""), "Estado", r["Estado"], "Validado", REGLA, fecha)
            P.poner(ws_c, fc, "Quechua_validado", x["final"])
            P.poner(ws_c, fc, "Estado", "Validado")
            validados += 1
        elif x["accion"].startswith("Revisar_Juan"):
            registrar(i, "Español (propuesta del traductor)", r["Español"], x["español_nuevo"],
                      REGLA + " — no se valida hasta que Juan decida")
            registrar(i, "Quechua (propuesta del traductor)", r["Quechua_validado"] or r["Quechua_borrador"], x["final"],
                      REGLA + " — corresponde al español propuesto")
            if r["Estado"] != "Revisar_Juan":
                agregar_cambio(ws_cam, i, codigos.get(i, ""), "Estado", r["Estado"], "Revisar_Juan", REGLA, fecha)
            P.poner(ws_c, fc, "Estado", "Revisar_Juan")
    ws_tm = wb_m["Traductor_metricas"]
    cols_tm = P.columnas(ws_tm)
    fila_dest = None
    for rr in range(2, ws_tm.max_row + 1):
        if ws_tm.cell(row=rr, column=cols_tm["Bloque"]).value == b:
            fila_dest = rr
    fila_dest = fila_dest or ws_tm.max_row + 1
    for k, v in fila_tm.items():
        if k in cols_tm:
            ws_tm.cell(row=fila_dest, column=cols_tm[k]).value = v
    P.guardar(wb_m, p_maestro)

    # ------------------------------------------------------------- escribir trabajo
    if p_trabajo.exists():
        wb_t = P.abrir(p_trabajo)
        ws_b, ws_d = wb_t["Bloque"], wb_t["Detalle"]
        fb, fd = P.filas_por_id(ws_b), P.filas_por_id(ws_d)
        for x in resultados:
            i = x["ID"]
            if x["accion"] == "Validado":
                P.poner(ws_b, fb[i], "Quechua", x["final"])
                P.poner(ws_b, fb[i], "Estado", "Validado")
                P.poner(ws_d, fd[i], "Quechua_validado", x["final"])
            elif x["accion"].startswith("Revisar_Juan"):
                P.poner(ws_b, fb[i], "Estado", "Revisar_Juan")
            if x["d"] is not None:
                P.poner(ws_d, fd[i], "%_edición", round(100 * x["d"], 1))
            com = x["comentario"]
            if x["español_nuevo"]:
                com = (f"El traductor propone otro español: «{x['español_nuevo']}» y este quechua: «{x['final']}». "
                       + com).strip()
            P.poner(ws_d, fd[i], "Comentario_traductor", com)
            P.poner(ws_d, fd[i], "Avisos_ortografía_traductor", "; ".join(x["orto"]))
        P.guardar(wb_t, p_trabajo)

    # ------------------------------------------------------------- reporte
    cl = Counter(x["clase"] for x in resultados)
    es_cambiado = [x for x in resultados if x["español_nuevo"]]
    md = [f"# Importación del traductor — {b}", "", f"Archivo leído: `{p_docx.name}` ({len(devuelto)} filas con ID en el Word).", ""]
    if n_com:
        md += [f"> **Atención:** el Word tiene {n_com} comentarios de Word (globos). No se importan: revísalos a mano.", ""]
    if sin_tocar:
        md += [f"> **Sublotes devueltos sin ningún cambio: {', '.join(sin_tocar)}.** Sus filas NO se validaron (puede que no se "
               "hayan revisado). Confirma con el traductor y, si los revisó, reimporta con "
               f"`--confirmar-sublotes {','.join(sin_tocar)}`.", ""]
    md += ["## Emparejamiento por ID", "",
           f"- Esperadas: {len(esperados)} · encontradas: {len(set(cuenta) & set(esperados))} · validadas: {validados}",
           f"- Faltantes: {', '.join(faltantes) or 'ninguna'}", f"- Sobrantes: {', '.join(sobrantes) or 'ninguna'}",
           f"- Duplicadas: {', '.join(duplicados) or 'ninguna'} (se usó la primera aparición)",
           f"- Sin quechua en el Word devuelto: {', '.join(x['ID'] for x in resultados if x['clase'] == 'vacío') or 'ninguna'}",
           f"- El traductor cambió el español en: {', '.join(x['ID'] for x in es_cambiado) or 'ninguna'} "
           "(no se validan: quedan Revisar_Juan; la propuesta está en Cambios y en el Excel de trabajo)", "",
           "## Métricas (Traductor_metricas)", ""]
    md += [f"- {k}: {v}" for k, v in fila_tm.items()]
    md += ["", "## Clases", ""] + [f"- {k}: {v}" for k, v in cl.most_common()]
    if es_cambiado:
        md += ["", "## Español cambiado por el traductor (decide Juan)", "", "| ID | Español actual | Propuesta | Quechua |", "|---|---|---|---|"]
        md += [f"| {x['ID']} | {corpus[x['ID']]['Español']} | {x['español_nuevo']} | {x['final']} |" for x in es_cambiado]
    md += ["", "## Comentarios del traductor", ""] + ([f"- `{x['ID']}`: {x['comentario']}" for x in resultados if x["comentario"]] or ["- ninguno"])
    md += ["", "## Avisos de ortografía sobre el quechua del traductor (solo informativo, no se corrige nada)", ""]
    md += [f"- `{x['ID']}`: {'; '.join(x['orto'])}" for x in resultados if x["orto"]] or ["- ninguno"]
    escribir_texto(p_reporte, "\n".join(md))
    print(f"OK {b}: {validados} validadas; faltantes {len(faltantes)}, sobrantes {len(sobrantes)}, duplicadas {len(duplicados)}; "
          f"español cambiado {len(es_cambiado)}; sublotes sin tocar {sin_tocar or 'ninguno'}; clases {dict(cl)}; reporte → {p_reporte}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
