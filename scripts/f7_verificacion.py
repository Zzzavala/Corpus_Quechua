"""F7 – Verificación final: checklist de CLAUDE.md §13, automático.

Revisa el maestro, los bloques con archivos, los Word y Excel (también abriéndolos con LibreOffice), corre las pruebas
(test_qa, test_ida_vuelta de 06–08), los reportes y que las fuentes sigan intactas.
Salida: 04_reportes/VERIFICACION_F7.md (PASA/FALLA por punto). Código 1 si algo falla.
Uso: python scripts/f7_verificacion.py [--sin-libreoffice] [--sin-pruebas]
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
from collections import Counter

import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook

from comun import CFG, RAIZ, cargar_checkpoint, clave_normalizada, escribir_texto, ruta
from heuristicas import tiempo_es

PY = sys.executable
R: list[tuple[str, str, bool, str]] = []
CB = CFG["bloque"]
MAX_PRESENTE = 0.65          # la parte general no puede ser casi toda presente (aclaración del proyecto, 14-oct)


def ok(seccion, nombre, cond, detalle=""):
    R.append((seccion, nombre, bool(cond), detalle))
    print(f"[{'PASA ' if cond else 'FALLA'}] {seccion} · {nombre} {detalle}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sin-libreoffice", action="store_true")
    ap.add_argument("--sin-pruebas", action="store_true", help="no correr test_qa ni test_ida_vuelta (usa el reporte anterior)")
    a = ap.parse_args()
    rep = ruta("reportes")

    # ------------------------------------------------------------ progreso y fuentes
    prog = (rep / "PROGRESO.md").read_text(encoding="utf-8")
    estados = dict(re.findall(r"^\| (F\d) – [^|]+\| (\w[\w ]*?) \|", prog, re.M))
    ok("Progreso", "PROGRESO.md con las fases F0–F7 en «hecha»", all(estados.get(f"F{i}") == "hecha" for i in range(8)), str(estados))
    r = subprocess.run([PY, "scripts/00_huella_fuentes.py", "--verificar"], cwd=RAIZ, capture_output=True, text=True, encoding="utf-8")
    ok("Fuentes", "Fuentes intactas (tamaño y sha256 iguales a F0)", r.returncode == 0, r.stdout.strip().replace("\n", " ")[:200])
    inv = (rep / "INVENTARIO.md").read_text(encoding="utf-8")
    ok("Fuentes", "Cifras de §5 verificadas (0 «✖ DIFIERE» en INVENTARIO.md)", "«✖ DIFIERE»: **0**" in inv)

    # ------------------------------------------------------------ maestro
    maestro = RAIZ / CFG["rutas"]["maestro"]
    H = pd.read_excel(maestro, sheet_name=None, dtype=str, keep_default_na=False)
    pestañas = ["LEEME", "Corpus", "Origen", "Cambios", "Descartes", "Lexico", "QA", "Bloques", "Traductor_metricas"]
    ok("Maestro", "corpus_maestro.xlsx con todas sus pestañas", all(p in H for p in pestañas), ", ".join(H))
    C = H["Corpus"]
    ok("Maestro", "IDs únicos y con formato GEN-/CLI-", C["ID"].is_unique and C["ID"].str.match(r"^(GEN|CLI)-\d{4}$").all(), f"({len(C)} filas)")
    ok("Maestro", "Toda fila de Corpus tiene su fila en Origen", set(C["ID"]) <= set(H["Origen"]["ID"]), "")
    cols = ["ID", "Dominio", "Escenario", "Etapa_consulta", "Rol", "Español", "Quechua_borrador", "Quechua_validado", "Estado",
            "Prioridad", "Bloque", "Sublote", "Split", "Test_sugerido"]
    ok("Maestro", "Corpus tiene las columnas de §8.2", list(C.columns[:len(cols)]) == cols, "")
    ok("Maestro", "Todo descarte tiene motivo", (H["Descartes"]["Motivo"] != "").all(), f"({len(H['Descartes'])} descartes)")
    ok("Maestro", "Cambios registrados", len(H["Cambios"]) > 0 and (H["Cambios"]["Regla"] != "").all(), f"({len(H['Cambios'])} cambios)")
    # trazabilidad completa: cada registro ingerido en F2 está en Corpus, Léxico o Descartes
    F2 = cargar_checkpoint("F2_ingesta")["Registros"]
    F5 = cargar_checkpoint("F5_huecos")
    uids_corpus = set(F5["Corpus"]["uid"])
    base_corpus = {u.split("/")[0] for u in uids_corpus}
    presentes = base_corpus | set(F5["Lexico"]["uid"]) | set(F5["Descartes"]["uid"])
    faltan = sorted(set(F2["uid"]) - presentes)
    ok("Maestro", "Nada se borra sin rastro (cada registro de F2 está en Corpus, Léxico o Descartes)", not faltan,
       f"({len(F2)} registros; sin rastro: {faltan[:5]})")

    # ------------------------------------------------------------ bloques con archivos
    F6 = cargar_checkpoint("F6_bloques")
    C6 = F6["Corpus"].set_index("ID")
    FB = F6["Filas_bloque"]
    ok("Bloques", "Ninguna fila repetida entre bloques", FB["ID"].is_unique, f"({len(FB)} filas en {FB['Bloque'].nunique()} bloques)")
    tol = CB["medicas_paciente_tolerancia"]
    for b in CB["bloques_con_archivos"]:
        f = FB[FB["Bloque"] == b]
        sub = C6.loc[f["ID"]]
        dom = Counter(sub["Dominio"])
        ok(b, "200 filas = 140 generales + 60 médicas", len(f) == CB["tamano"] and dom["General"] == CB["generales"] and dom["Médico"] == CB["medicas"],
           f"({len(f)}: {dom['General']} + {dom['Médico']})")
        tam = Counter(zip(sub["Sublote"], sub["Dominio"]))
        subl = sorted(set(sub["Sublote"]))
        bien = len(subl) == CB["sublotes"] and all(tam[(s, "General")] + tam[(s, "Médico")] == CB["tamano_sublote"] and
                                                   abs(tam[(s, "Médico")] - CB["medicas_por_sublote"]) <= 1 for s in subl)
        ok(b, "4 sublotes de 50 con ~35 + ~15", bien, " | ".join(f"{s}: {tam[(s, 'General')]}+{tam[(s, 'Médico')]}" for s in subl))
        roles = Counter(sub.loc[sub["Dominio"] == "Médico", "Rol"])
        ok(b, "Médico / paciente ≈ 50/50", abs(roles["Paciente"] - CB["medicas_paciente_objetivo"]) <= tol and roles["Médico"] + roles["Paciente"] == CB["medicas"],
           f"({roles['Médico']} / {roles['Paciente']})")
        resp = sub[sub["Responde_a"] != ""]
        pares_ok = all(q in sub.index and sub.loc[q, "Sublote"] == r["Sublote"] for q, (_, r) in zip(resp["Responde_a"], resp.iterrows()))
        ok(b, "Pares pregunta–respuesta en el mismo sublote", pares_ok and len(resp) > 0, f"({len(resp)} respuestas)")
        if b == "B01":
            etapas = set(sub["Plan_tema"]) - {""}
            esperadas = {e["etapa"] for e in CFG["medico"]["b01_etapas"]}
            ok(b, "Consulta completa: las 6 etapas de §11.2", esperadas <= etapas, ", ".join(sorted(etapas)))
            ctrl = sub[sub["Control"] == "Sí"]
            ok(b, "~15 oraciones de control (10 generales + 5 médicas), en los 4 sublotes",
               len(ctrl) == 15 and Counter(ctrl["Dominio"]) == Counter({"General": 10, "Médico": 5}) and ctrl["Sublote"].nunique() == 4,
               f"({len(ctrl)}: {dict(Counter(ctrl['Dominio']))})")

        # Word del traductor
        doc = Document(ruta("bloques") / b / f"{b}_traductor.docx")
        sec = doc.sections[0]
        ok(b, "Word: A4 horizontal", sec.page_width > sec.page_height and abs(sec.page_width.cm - 29.7) < 0.1, f"({sec.page_width.cm:.1f} × {sec.page_height.cm:.1f} cm)")
        texto = "\n".join(p.text for p in doc.paragraphs)
        ok(b, "Word: instrucciones y convención ortográfica en la página 1", "Qué hacer" in texto and "Convención ortográfica" in texto and "No borre" in texto)
        tablas = [t for t in doc.tables if t.rows[0].cells[0].text.strip() == "N°"]
        ok(b, "Word: una tabla por sublote", len(tablas) == CB["sublotes"], f"({len(tablas)})")
        rep_enc = all(t.rows[0]._tr.trPr is not None and t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None for t in tablas)
        ok(b, "Word: fila de encabezado repetida en cada página", rep_enc)
        enc = [c.text for c in tablas[0].rows[0].cells] if tablas else []
        ok(b, "Word: columnas N° | ID | Quién habla | Español | Quechua | Observaciones",
           enc == ["N°", "ID", "Quién habla", "Español", "Quechua (corrija o escriba aquí)", "Observaciones"], str(enc))
        filas_w = [(row.cells[1].text, row.cells[4].text, row.cells[5].text) for t in tablas for row in t.rows[1:]]
        ok(b, "Word: mismas 200 filas y en el mismo orden (por ID)", [x[0] for x in filas_w] == list(f["ID"]), f"({len(filas_w)})")
        todo = texto + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        prohibidos = ["Lopera", "Rosettia", "SOCEMURP", "Allin Kawsay", "MINEDU", "OPS/OMS", "Farfán", "unificado", "traductor automático",
                      "Tipo_borrador", "Estado_unificado", "Por_revisar", "Fuente"]
        hallados = [p for p in prohibidos if p in todo]
        ok(b, "Word: sin información de origen ni columnas técnicas", not hallados, str(hallados))
        mal_qu = [i for i, q, _ in filas_w if q != ("" if C6.loc[i, "Control"] == "Sí" else C6.loc[i, "Quechua_borrador"])]
        ok(b, "Word: celda Quechua = borrador; vacía en control y en filas sin borrador", not mal_qu, f"(distintas: {mal_qu[:5]})")
        obs_mal = [i for i, _, o in filas_w if o and not o.startswith("Aviso:")]
        ok(b, "Word: Observaciones solo con «Aviso:» (ALTA) o vacía", not obs_mal, f"({sum(1 for x in filas_w if x[2])} avisos)")

        # Excel de trabajo
        wb = load_workbook(ruta("bloques") / b / f"{b}_trabajo.xlsx")
        ok(b, "Excel de trabajo: pestañas Bloque / Detalle / Resumen (Bloque activa)", wb.sheetnames == ["Bloque", "Detalle", "Resumen"] and wb.active.title == "Bloque",
           str(wb.sheetnames))
        ws = wb["Bloque"]
        enc_b = [c.value for c in ws[1]]
        ok(b, "Excel de trabajo: columnas de §8.3", enc_b == ["N°", "ID", "Sublote", "Dominio", "Rol", "Español", "Quechua", "Estado", "Grabado_ES", "Grabado_QU", "Notas"], str(enc_b))
        ok(b, "Excel de trabajo: encabezado inmovilizado, autofiltro y lista en Estado",
           ws.freeze_panes == "A2" and ws.auto_filter.ref and len(ws.data_validations.dataValidation) >= 1, "")
        ok(b, "Excel de trabajo: solo valores (sin fórmulas)",
           not any(isinstance(c.value, str) and c.value.startswith("=") for hoja in wb for fila in hoja.iter_rows() for c in fila), "")

    # ------------------------------------------------------------ contenido entre bloques
    claves = FB["ID"].map(lambda i: clave_normalizada(C6.loc[i, "Español"]))
    rep_es = claves[claves.duplicated(keep=False)]
    ok("Bloques", "Ningún español repetido entre las filas de los bloques (§10.4)", rep_es.empty,
       f"(repetidos: {sorted(set(FB.loc[rep_es.index, 'ID']))[:6]})")
    for b in CB["bloques_con_archivos"]:
        gen = [i for i in FB.loc[FB["Bloque"] == b, "ID"] if C6.loc[i, "Dominio"] == "General"]
        t = Counter(tiempo_es(C6.loc[i, "Español"]) for i in gen)
        pres = t.get("presente", 0) / max(1, len(gen))
        ok(b, f"Parte general con varios tiempos (presente ≤ {100 * MAX_PRESENTE:.0f} %; requisito del proyecto)",
           pres <= MAX_PRESENTE and len(t) >= 5, " · ".join(f"{k} {v}" for k, v in t.most_common()))

    # ------------------------------------------------------------ scripts, guía, reportes
    if a.sin_pruebas:
        prueba = (rep / "PRUEBA_06_07_08.md").read_text(encoding="utf-8")
        ok("Scripts", "06, 07 y 08 probados con datos simulados (reporte anterior; se usó --sin-pruebas)", "TODO PASA" in prueba,
           re.search(r"\((\d+ de \d+)\)", prueba).group(1) if re.search(r"\((\d+ de \d+)\)", prueba) else "")
    else:
        for script, nombre in [("test_qa.py", "Pruebas de la revisión automática y la limpieza (test_qa)"),
                               ("test_ida_vuelta.py", "06, 07 y 08 probados ahora con datos simulados (ida y vuelta por ID)")]:
            r = subprocess.run([PY, f"scripts/{script}"], cwd=RAIZ, capture_output=True, text=True, encoding="utf-8", timeout=1800)
            fallas = [x for x in (r.stdout or "").splitlines() if "FALLA" in x]
            ok("Scripts", nombre, r.returncode == 0, f"({len(fallas)} fallas{': ' + fallas[0][:120] if fallas else ''})")
    ok("Scripts", "Existen 06_importar_traductor.py, 07_hoja_grabacion.py y 08_manifest.py",
       all((RAIZ / "scripts" / s).exists() for s in ["06_importar_traductor.py", "07_hoja_grabacion.py", "08_manifest.py"]))
    ok("Grabación", "GUIA_GRABACION.docx creada", (ruta("bloques") / "GUIA_GRABACION.docx").exists())
    eda = (rep / "EDA.md").read_text(encoding="utf-8")
    ok("Reportes", "EDA.md con conclusiones en lenguaje simple", "## 11. Conclusiones" in eda)
    dec = (rep / "DECISIONES_TOMADAS.md").read_text(encoding="utf-8")
    ok("Reportes", "DECISIONES_TOMADAS.md con los valores por defecto P1–P7 y las decisiones de cada fase",
       all(f"| P{i} |" in dec for i in range(1, 8)) and all(f"## F{i}" in dec for i in range(0, 7)))
    ent = rep / "ENTREGA.md"
    if ent.exists():
        t = ent.read_text(encoding="utf-8")
        claves = ["Qué está listo", "Qué enviar a quién", "Qué debes revisar", "Pendientes", "Cuántos bloques"]
        ok("Reportes", "ENTREGA.md con: listo, qué enviar a quién, qué revisar, pendientes P1–P7, cuántos bloques",
           all(k in t for k in claves) and all(f"P{i}" in t for i in range(1, 8)), str([k for k in claves if k not in t]))
    else:
        ok("Reportes", "ENTREGA.md existe", False)

    # ------------------------------------------------------------ abrir de verdad con LibreOffice
    if not a.sin_libreoffice:
        archivos = [ruta("bloques") / b / f"{b}_{x}" for b in CB["bloques_con_archivos"] for x in ("traductor.docx", "trabajo.xlsx")]
        archivos += [ruta("bloques") / "GUIA_GRABACION.docx", maestro]
        for p in archivos:
            r = subprocess.run([PY, "scripts/render_revision.py", str(p), "--paginas", "1"], cwd=RAIZ, capture_output=True, text=True,
                               encoding="utf-8", timeout=1200)
            m = re.search(r"\((\d+) páginas\)", r.stdout or "")
            ok("LibreOffice", f"{p.name} abre y se convierte a PDF", r.returncode == 0 and m, f"({m.group(1)} páginas)" if m else (r.stderr or "")[-150:])

    total = len(R)
    pasan = sum(1 for x in R if x[2])
    md = ["# Verificación final (F7) — checklist de CLAUDE.md §13", "",
          f"Generado por `scripts/f7_verificacion.py`. **{pasan} de {total} puntos PASAN.**", "",
          "| Sección | Punto | Resultado | Detalle |", "|---|---|---|---|"]
    md += [f"| {s} | {n} | {'PASA' if o else '**FALLA**'} | {d.replace('|', '/')} |" for s, n, o, d in R]
    if a.sin_libreoffice:
        md += ["", "Esta corrida usó `--sin-libreoffice`: los archivos NO se abrieron con LibreOffice."]
    else:
        md += ["", "Las páginas convertidas por LibreOffice están en `04_reportes/revision_visual/` (PDF + PNG de la primera "
               "página de cada archivo)."]
    if a.sin_pruebas:
        md += ["", "Esta corrida usó `--sin-pruebas`: el punto de 06–08 se tomó del reporte anterior `PRUEBA_06_07_08.md`."]
    escribir_texto(rep / "VERIFICACION_F7.md", "\n".join(md))
    print(f"\n{pasan}/{total} PASAN → 04_reportes/VERIFICACION_F7.md")
    return 0 if pasan == total else 1


if __name__ == "__main__":
    sys.exit(main())
