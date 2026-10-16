"""Lectura de las fuentes (solo lectura). Cada función devuelve un DataFrame con la fila original.

Nunca se escribe sobre las fuentes. Los Excel con fórmulas se leen con data_only=True y los IDs
que dependen de fórmulas (N° de Grupo 3 y de Corpus_Construido) se reconstruyen por posición.
"""
from __future__ import annotations

import html as html_mod
import json
import re
from functools import lru_cache
from pathlib import Path

import openpyxl
import pandas as pd

from comun import CFG, RAIZ, asegurar_dir, escribir_json, fuente, ruta, texto

# --------------------------------------------------------------------------- Excel genérico


def _filas(path: Path, hoja: str, data_only: bool = True) -> list[tuple[int, list]]:
    wb = openpyxl.load_workbook(path, data_only=data_only, read_only=True)
    ws = wb[hoja]
    out = []
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        out.append((i, list(r)))
    wb.close()
    return out


def hojas(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    n = list(wb.sheetnames)
    wb.close()
    return n


def resumen_hoja(path: Path, hoja: str) -> dict:
    filas = _filas(path, hoja)
    no_vacias = [f for f in filas if any(not _vacio(c) for c in f[1])]
    ncol = max((len(f[1]) for f in filas), default=0)
    encabezado = [texto(c) for c in filas[0][1]] if filas else []
    while encabezado and encabezado[-1] == "":
        encabezado.pop()
    return {"hoja": hoja, "filas_totales": len(filas), "filas_no_vacias": len(no_vacias), "columnas": ncol,
            "encabezado": encabezado}


def _vacio(c) -> bool:
    return c is None or (isinstance(c, str) and c.strip() == "")


def _col(r: list, i: int):
    return r[i] if i < len(r) else None


# --------------------------------------------------------------------------- unificado (§5.1)

COLS_UNIFICADO = ["Id_unificado", "Español (palabra o frase)", "Quechua_final (chanka)", "quechua_marshall",
                  "quechua_kim1", "quechua_kim2", "Observaciones_finales", "Estado", "Sección",
                  "Fuente_original", "Códigos_origen"]


@lru_cache(maxsize=1)
def cargar_unificado() -> pd.DataFrame:
    filas = _filas(fuente("unificado"), "Vocabulario_unificado")
    enc = [texto(c) for c in filas[0][1]]
    idx = {h: enc.index(h) for h in COLS_UNIFICADO}
    datos = []
    for n, r in filas[1:]:
        if all(_vacio(c) for c in r):
            continue
        d = {h: texto(_col(r, idx[h])) for h in COLS_UNIFICADO}
        d["Fila_original"] = n
        datos.append(d)
    return pd.DataFrame(datos)


@lru_cache(maxsize=1)
def cargar_trazabilidad() -> pd.DataFrame:
    filas = _filas(fuente("unificado"), "Trazabilidad")
    enc = [texto(c) for c in filas[0][1]]
    datos = []
    for n, r in filas[1:]:
        if all(_vacio(c) for c in r):
            continue
        d = {h: texto(_col(r, i)) for i, h in enumerate(enc) if h}
        d["Fila_original"] = n
        datos.append(d)
    return pd.DataFrame(datos)


def fuente_canonica(fuente_original: str) -> str:
    f = texto(fuente_original)
    reglas = [
        ("Lopera", "Lopera Quiroga (2020)"),
        ("Allin Kawsay", "Allin Kawsay (INS/MINSA)"),
        ("OPS/OMS Bolivia", "OPS/OMS Bolivia (2013)"),
        ("Corpus original del equipo", "SOCEMURP (corpus original del equipo)"),
        ("Manual de quechua médico SOCEMURP", "SOCEMURP (corpus original del equipo)"),
        ("Pregunta añadida por Kim 1", "Kim 1 (pregunta añadida)"),
        ("Farfán", "Farfán Cantero (2001), SITUA"),
        ("Manual del quechua chanka para la administración de justicia", "MinCul – manual de justicia (chanka)"),
        ("YouTube", "YouTube"),
        ("Afiche MINSA/MinCul", "Afiche MINSA/MinCul lavado de manos (chanka)"),
        ("MINSA", "MINSA – diabetes (quechua)"),
        ("OMS (2009)", "OMS (2009) – lista de verificación"),
    ]
    for clave, nombre in reglas:
        if f.startswith(clave):
            return nombre
    return f or "(sin fuente)"


# --------------------------------------------------------------------------- Corpus_Equipos (§5.2)


def _hoja_simple(hoja: str, col_es: int, col_qu: int, col_obs: int, col_id: int = 0, col_tipo: int | None = None):
    filas = _filas(fuente("corpus_equipos"), hoja)
    datos = []
    for n, r in filas[1:]:
        if all(_vacio(c) for c in r):
            continue
        d = {
            "Código_original": texto(_col(r, col_id)),
            "Español": texto(_col(r, col_es)),
            "Quechua": texto(_col(r, col_qu)),
            "Observaciones": texto(_col(r, col_obs)),
            "Extra": " | ".join(texto(c) for c in r[max(col_obs, col_qu) + 1:] if not _vacio(c)),
            "Fila_original": n,
        }
        if col_tipo is not None:
            d["Tipo"] = texto(_col(r, col_tipo))
        datos.append(d)
    return pd.DataFrame(datos)


@lru_cache(maxsize=1)
def cargar_principito() -> pd.DataFrame:
    return _hoja_simple("Grupo 1 - Principito", 1, 2, 3)


@lru_cache(maxsize=1)
def cargar_grupo2() -> pd.DataFrame:
    return _hoja_simple("Grupo 2 - Corpus Chanka", 1, 2, 3)


@lru_cache(maxsize=1)
def cargar_grupo3() -> pd.DataFrame:
    df = _hoja_simple("Grupo 3 - Corpus Medico", 1, 2, 3)
    # N° es fórmula (=A2+1): se reconstruye por posición.
    df["Código_original"] = [f"G3-{i:04d}" for i in range(1, len(df) + 1)]
    return df


@lru_cache(maxsize=1)
def cargar_terminos() -> pd.DataFrame:
    df = _hoja_simple("Grupo 3 - Términos médicos", 2, 3, 4, col_tipo=1)
    df["Código_original"] = [f"G3T-{i:03d}" for i in range(1, len(df) + 1)]
    return df


@lru_cache(maxsize=1)
def cargar_construido() -> pd.DataFrame:
    filas = _filas(fuente("corpus_construido"), "Hoja 1")
    datos = []
    for n, r in filas[1:]:
        if all(_vacio(c) for c in r):
            continue
        datos.append({"Español": texto(_col(r, 2)), "Quechua": texto(_col(r, 3)), "Fila_original": n})
    df = pd.DataFrame(datos)
    # N° es fórmula (=CONCATENATE("S1-",B2)): se reconstruye.
    df["Código_original"] = [f"S1-{i}" for i in range(1, len(df) + 1)]
    return df


# --------------------------------------------------------------------------- AmericasNLP (§5.5)


def _lineas(p: Path) -> list[str]:
    return p.read_text(encoding="utf-8-sig").splitlines()


@lru_cache(maxsize=1)
def americasnlp() -> dict[str, dict]:
    base = RAIZ / CFG["rutas"]["fuentes"] / CFG["archivos_fuente"]["americasnlp_dir"]
    out = {}
    for nombre, es, qu in [
        ("train (JW300)", "train.es", "train.quy"),
        ("dev (sureño estándar)", "dev.es", "dev.quy"),
        ("dict (MINEDU + Dict_misc)", "dict.es", "dict.quy"),
        ("es-quy/jw300", "es-quy/jw300.es-quy.es", "es-quy/jw300.es-quy.quy"),
        ("es-quy/minedu", "es-quy/minedu.quy-es.es", "es-quy/minedu.quy-es.quy"),
        ("es-quy/dict_misc", "es-quy/dict_misc.quy-es.es", "es-quy/dict_misc.quy-es.quy"),
    ]:
        pe, pq = base / es, base / qu
        if pe.exists() and pq.exists():
            out[nombre] = {"es": _lineas(pe), "qu": _lineas(pq), "archivos": [es, qu]}
    pt = base / "test.es"
    if pt.exists():
        out["test (solo español)"] = {"es": _lineas(pt), "qu": [], "archivos": ["test.es"]}
    return out


# --------------------------------------------------------------------------- PDF y HTML


def _cache_dir() -> Path:
    return asegurar_dir(ruta("checkpoints") / "cache")


def texto_pdf(clave: str) -> list[str]:
    """Texto por página (vacío si el PDF es solo imagen). Se guarda en caché."""
    p = fuente(clave)
    cache = _cache_dir() / f"pdf_{clave}.json"
    if cache.exists():
        datos = json.loads(cache.read_text(encoding="utf-8"))
        if datos.get("bytes") == p.stat().st_size:
            return datos["paginas"]
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(p)
    paginas = []
    for i in range(len(doc)):
        tp = doc[i].get_textpage()
        paginas.append(tp.get_text_range())
    doc.close()
    escribir_json(cache, {"bytes": p.stat().st_size, "paginas": paginas})
    return paginas


def ejemplos_minedu_pdf() -> list[dict]:
    """Ejemplos (quechua + traducción) del diccionario MINEDU en PDF.

    Formato observado: líneas que empiezan con '- ' (o '1. ', '2. ' en acepciones); un ejemplo quechua
    seguido de su traducción al español. Las líneas de continuación se pegan a la anterior.
    """
    paginas = texto_pdf("minedu_pdf")
    items = []
    for npag, t in enumerate(paginas, start=1):
        bloques: list[str] = []
        actual = None
        for linea in t.replace("\r", "\n").split("\n"):
            s = linea.strip()
            if not s or "indd" in s or re.fullmatch(r"\d{1,3}", s):
                continue
            m = re.match(r"^(?:-|\d\.)\s+(.*)$", s)
            if m:
                if actual is not None:
                    bloques.append(actual)
                actual = m.group(1)
            elif actual is not None and not re.match(r"^[A-ZÑ][a-zñ' ]+\.?\s*(\d\.)?\s*\((s|r|rt|ch|tp|sy)\)", s):
                actual += " " + s
            else:
                if actual is not None:
                    bloques.append(actual)
                actual = None
        if actual is not None:
            bloques.append(actual)
        for b in bloques:
            items.append({"pagina": npag, "texto": re.sub(r"\s+", " ", b).strip()})
    return items


def texto_html_farfan() -> str:
    raw = fuente("web_html").read_bytes()
    try:
        h = raw.decode("utf-8")
        if "�" in h or "Ã" in h:
            raise UnicodeDecodeError("utf-8", b"", 0, 1, "mojibake")
    except UnicodeDecodeError:
        h = raw.decode("cp1252", errors="replace")
    t = re.sub(r"<script.*?</script>|<style.*?</style>", " ", h, flags=re.S | re.I)
    t = re.sub(r"<[^>]+>", " ", t)
    return html_mod.unescape(re.sub(r"\s+", " ", t)).strip()
