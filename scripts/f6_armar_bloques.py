"""F6 – Priorización y bloques (§11) + archivos de los bloques de config.bloque.bloques_con_archivos (§8).

- Parte general (140/bloque): selección voraz por cobertura de sufijos, bigramas y palabras (§11.1), con cuotas por pool.
- Parte médica (60/bloque): plan de F5 (B01 curado; B02… por tema), unidades pregunta+respuestas indivisibles.
- 4 sublotes de 50 (≈35 + 15); control en B01 (D7); Test_sugerido 10 % estratificado (P1).
- Escribe 02_maestro/corpus_maestro.xlsx, 03_bloques/BXX/BXX_trabajo.xlsx y BXX_traductor.docx (bloques con archivos),
  04_reportes/REVISION_JUAN.md y el checkpoint F6.

Si el maestro ya tiene quechua validado (devoluciones del traductor), se conserva al regenerar.
Uso: python scripts/f6_armar_bloques.py [--forzar]
"""
from __future__ import annotations

import argparse
import random
import re
import sys
from collections import Counter, defaultdict

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from rapidfuzz import fuzz
from rapidfuzz.process import cdist

from comun import (CFG, RAIZ, asegurar_dir, cargar_checkpoint, clave_normalizada, es_pregunta, escribir_json, escribir_texto,
                   guardar_excel, guardar_seguro, hoy, n_palabras, ruta, ruta_checkpoint)
import docs_word as W
from heuristicas import tiempo_es
from excel_formato import escribir_clave_valor, escribir_hoja
from qa import aviso_word
import quechua as Q

CK = ruta_checkpoint("F6_bloques")
MAESTRO = RAIZ / CFG["rutas"]["maestro"]
BLOQUES_DIR = ruta("bloques")
CB = CFG["bloque"]
POOLS = list(CFG["general"]["orden_pools"])          # orden = desempate en la selección voraz
POOL_CLAUDE = "g_claude"
NOMBRE_POOL = {"principito": "Principito", "g2_diccionario": "G2 frases variadas", "g2_ejercicio": "G2 ejercicios",
               POOL_CLAUDE: "Redactadas por Claude (tiempos)"}
COLS_CORPUS = ["ID", "Dominio", "Escenario", "Etapa_consulta", "Rol", "Español", "Quechua_borrador", "Quechua_validado",
               "Estado", "Prioridad", "Bloque", "Sublote", "Split", "Test_sugerido", "Control", "Responde_a", "N_bloque"]
COLS_ORIGEN = ["ID", "Fuente", "Archivo", "Hoja", "Código_original", "Fila_original", "Tipo_borrador", "Variedad_fuente",
               "Origen_verificado", "Estado_unificado", "Observaciones_heredadas", "Códigos_origen", "Códigos_unidos",
               "Otros_borradores", "Propuestas_equipo", "Parte_de", "Evidencia_P6", "Fuente_detalle"]


# ======================================================================= selección general (§11.1)

def seleccionar_general(C: pd.DataFrame, bloques: list[str], excluir: set[str] = frozenset()) -> tuple[dict[str, list[str]], list[str]]:
    G = C[(C["Dominio"] == "General") & (~C["ID"].isin(excluir))].copy()
    G["_fila"] = G["Fila_original"].astype(int)
    G["_hoja"] = G["Hoja"].map({"Grupo 1 - Principito": 0, "Grupo 2 - Corpus Chanka": 1})
    G = G.sort_values(["_hoja", "_fila", "ID"])
    unidades = {r["ID"]: Q.unidades_cobertura(r["Quechua_borrador"]) for _, r in G.iterrows()}
    largo_qu = {r["ID"]: max(1, len(Q.tokens_qu(r["Quechua_borrador"]))) for _, r in G.iterrows()}
    largo_es = {r["ID"]: n_palabras(r["Español"]) for _, r in G.iterrows()}
    pool_de = dict(zip(G["ID"], G["Pool"]))
    restantes = {p: list(G.loc[G["Pool"] == p, "ID"]) for p in POOLS}
    restantes[POOL_CLAUDE] = orden_redactadas(G)
    pesos = CFG["general"]["pesos_cobertura"]
    expo = CFG["general"]["penalizacion_longitud_exponente"]
    cuotas_base = CFG["general"]["cuotas_por_bloque"]
    notas = []
    salida = {}
    for b in bloques:
        disp = {p: len(restantes[p]) for p in POOLS + [POOL_CLAUDE]}
        cuota = {"principito": min(cuotas_base["principito"], disp["principito"]),
                 "g2_diccionario": min(cuotas_base["g2_diccionario"], disp["g2_diccionario"]),
                 POOL_CLAUDE: min(cuotas_base["g_claude_min"], disp[POOL_CLAUDE])}
        cuota["g2_ejercicio"] = max(0, min(cuotas_base["g2_ejercicio_max"], disp["g2_ejercicio"],
                                           CB["generales"] - sum(cuota.values())))
        # Si las cuotas no llegan a 140, se completa en el orden de config (por defecto: más Principito —prioridad del
        # profesor—, frases variadas, oraciones de Claude en varios tiempos y, solo al final, ejercicios sobre el tope).
        falta = CB["generales"] - sum(cuota.values())
        for p in CFG["general"]["orden_relleno"]:
            extra = max(0, min(falta, disp[p] - cuota[p]))
            cuota[p] += extra
            falta -= extra
        if sum(cuota.values()) < CB["generales"]:
            notas.append(f"{b}: solo hay {sum(cuota.values())} oraciones generales (faltan {CB['generales'] - sum(cuota.values())}).")
            break
        if cuota["principito"] > cuotas_base["principito"]:
            notas.append(f"{b}: {cuota['principito']} frases del Principito (más que ~{cuotas_base['principito']}) porque solo quedaban "
                         f"{disp['g2_diccionario']} frases variadas y los ejercicios tienen tope de {cuotas_base['g2_ejercicio_max']}.")
        if cuota[POOL_CLAUDE] > cuotas_base["g_claude_min"]:
            notas.append(f"{b}: {cuota[POOL_CLAUDE]} oraciones generales redactadas por Claude (solo español, varios tiempos; mínimo "
                         f"{cuotas_base['g_claude_min']}) porque solo quedaban {disp['principito']} frases del Principito y "
                         f"{disp['g2_diccionario']} frases variadas, y los ejercicios tienen tope de {cuotas_base['g2_ejercicio_max']}.")
        if cuota["g2_ejercicio"] > cuotas_base["g2_ejercicio_max"]:
            notas.append(f"{b}: {cuota['g2_ejercicio']} ejercicios del Grupo 2 (más que el tope de {cuotas_base['g2_ejercicio_max']}) porque "
                         f"no quedaban otras oraciones generales.")
        max_pal = CFG["general"]["b01_max_palabras_es"] if b == "B01" else CFG["general"]["max_palabras_es"]
        cubiertas = {"sufijo": set(), "bigrama": set(), "palabra": set()}
        elegidos = []
        restante = dict(cuota)
        candidatos = {p: [i for i in restantes[p] if largo_es[i] <= max_pal] for p in POOLS}
        for p in POOLS:                                     # si no alcanzan las cortas, se relaja el largo
            if len(candidatos[p]) < cuota[p]:
                candidatos[p] = list(restantes[p])
        while len(elegidos) < CB["generales"] - cuota[POOL_CLAUDE]:
            mejor, mejor_clave = None, None
            for p in POOLS:
                if restante[p] <= 0:
                    continue
                for i in candidatos[p]:
                    u = unidades[i]
                    ganancia = sum(pesos[t] * len(u[t] - cubiertas[t]) for t in pesos)
                    puntaje = ganancia / (largo_qu[i] ** expo)
                    clave = (puntaje, -POOLS.index(p), -largo_es[i])
                    if mejor_clave is None or clave > mejor_clave:
                        mejor, mejor_clave = i, clave
            p = pool_de[mejor]
            elegidos.append(mejor)
            restante[p] -= 1
            candidatos[p].remove(mejor)
            restantes[p].remove(mejor)
            for t in cubiertas:
                cubiertas[t] |= unidades[mejor][t]
        # las redactadas por Claude no tienen borrador (no hay sufijos que medir): se toman en su orden intercalado por
        # rasgo gramatical, para que cada bloque tenga una mezcla pareja de tiempos
        redactadas = restantes[POOL_CLAUDE][:cuota[POOL_CLAUDE]]
        del restantes[POOL_CLAUDE][:cuota[POOL_CLAUDE]]
        salida[b] = intercalar(elegidos, redactadas)
    return salida, notas


def resumen_tiempos(ids: list[str], C: pd.DataFrame) -> str:
    """Mezcla de tiempos medida sobre el verbo en ESPAÑOL, igual para todas las filas (fuentes y redactadas)."""
    info = C.set_index("ID")
    c = Counter(tiempo_es(info.loc[i, "Español"]) for i in ids)
    n = max(1, len(ids))
    preg = sum(1 for i in ids if es_pregunta(info.loc[i, "Español"]))
    return " · ".join(f"{k} {v} ({100 * v / n:.0f} %)" for k, v in c.most_common()) + f" · (preguntas: {preg})"


def orden_redactadas(G: pd.DataFrame) -> list[str]:
    """Orden de las oraciones generales de Claude: cada rasgo repartido en proporción a su tamaño (mezcla pareja)."""
    R = G[G["Pool"] == POOL_CLAUDE]
    por_cat: dict[str, list[tuple[int, str]]] = {}
    for u, i, k in zip(R["uid"], R["ID"], R["Fila_original"]):
        por_cat.setdefault(u.split(":")[2], []).append((int(k), i))          # Fila_original = posición en el YAML
    claves = []
    for ci, (cat, lst) in enumerate(por_cat.items()):
        lst.sort()
        for j, (_, i) in enumerate(lst):
            claves.append(((j + 0.5) / len(lst), ci, i))
    return [i for _, _, i in sorted(claves)]


def intercalar(a: list[str], b: list[str]) -> list[str]:
    """Mezcla dos listas manteniendo su orden y repartiendo b de forma pareja entre a."""
    claves = [((k + 0.5) / len(a), 0, x) for k, x in enumerate(a)] + [((k + 0.5) / len(b), 1, x) for k, x in enumerate(b)]
    return [x for _, _, x in sorted(claves)]


# ======================================================================= armado de bloques


def armar_bloque(b: str, gen_ids: list[str], C: pd.DataFrame) -> list[dict]:
    """Devuelve las 200 filas del bloque en orden, con N° y Sublote."""
    M = C[C["Plan_bloque"] == b].copy()
    M["Plan_orden"] = M["Plan_orden"].astype(int)
    M = M.sort_values("Plan_orden")
    unidades: list[list[str]] = []
    actual = None
    for _, r in M.iterrows():
        if r["Plan_unidad"] != actual:
            unidades.append([])
            actual = r["Plan_unidad"]
        unidades[-1].append(r["ID"])
    if sum(len(u) for u in unidades) != CB["medicas"]:
        raise ValueError(f"{b}: la parte médica tiene {sum(len(u) for u in unidades)} filas, no {CB['medicas']}")
    # unidades → sublotes por posición de inicio (15 por sublote)
    sub_med = defaultdict(list)
    pos = 0
    for u in unidades:
        s = min(CB["sublotes"] - 1, pos // CB["medicas_por_sublote"])
        sub_med[s].append(u)
        pos += len(u)
    m_size = {s: sum(len(u) for u in sub_med[s]) for s in range(CB["sublotes"])}
    g_cap = {s: CB["tamano_sublote"] - m_size[s] for s in range(CB["sublotes"])}
    # Generales repartidas por ORIGEN: cada origen (Principito, frases variadas, ejercicios, redactadas) se reparte en
    # ronda entre los sublotes, para que ningún sublote se quede con todos los ejercicios (§11.4). Se respeta la capacidad.
    pool_de = dict(zip(C["ID"], C["Pool"]))
    por_pool: dict[str, list[str]] = defaultdict(list)
    for i in gen_ids:
        por_pool[pool_de[i]].append(i)
    sub_gen = defaultdict(list)
    s = 0
    for p in sorted(por_pool, key=lambda p: -len(por_pool[p])):
        for i in por_pool[p]:
            intentos = 0
            while len(sub_gen[s]) >= g_cap[s] and intentos < CB["sublotes"]:
                s = (s + 1) % CB["sublotes"]
                intentos += 1
            sub_gen[s].append(i)
            s = (s + 1) % CB["sublotes"]
    # dentro de cada sublote, mezclar los orígenes (no todos los ejercicios seguidos)
    for k in sub_gen:
        grupos = defaultdict(list)
        for i in sub_gen[k]:
            grupos[pool_de[i]].append(i)
        claves = [((j + 0.5) / len(lst), p, i) for p, lst in grupos.items() for j, i in enumerate(lst)]
        sub_gen[k] = [i for _, _, i in sorted(claves)]
    filas = []
    n = 0
    for s in range(CB["sublotes"]):
        us = sub_med[s]
        gs = sub_gen[s]
        trozos = len(us) + 1
        base, resto = divmod(len(gs), trozos)
        tam = [base + (1 if k < resto else 0) for k in range(trozos)]
        k0 = 0
        orden_ids = []
        for k in range(trozos):
            orden_ids += [("G", i) for i in gs[k0:k0 + tam[k]]]
            k0 += tam[k]
            if k < len(us):
                orden_ids += [("M", i) for i in us[k]]
        for _, i in orden_ids:
            n += 1
            filas.append({"ID": i, "N°": n, "Sublote": f"{b}-S{s + 1}"})
    return filas


def elegir_control(filas_b01: list[dict], C: pd.DataFrame, visibles: list[str]) -> set[str]:
    """D7: ~10 generales + ~5 médicas de B01 con borrador, repartidas entre sublotes.
    No se elige una fila que tenga una «hermana» con borrador visible muy parecido en la misma entrega (B01+B02):
    el traductor vería casi la misma solución y el análisis de anclaje saldría sesgado."""
    if not CFG["control"]["activo"]:
        return set()
    rng = random.Random(CFG["semilla"])
    info = C.set_index("ID")
    umbral = CFG["control"]["similitud_max_hermana"]
    excluir = set(CFG["control"].get("excluir_codigos") or [])      # el traductor solo verá el español: debe ser claro
    vis = [i for i in visibles if info.loc[i, "Quechua_borrador"]]
    vis_qu = [clave_normalizada(info.loc[i, "Quechua_borrador"]) for i in vis]
    vis_es = [clave_normalizada(info.loc[i, "Español"]) for i in vis]

    def tiene_hermana(i: str) -> bool:
        q, e = clave_normalizada(info.loc[i, "Quechua_borrador"]), clave_normalizada(info.loc[i, "Español"])
        for j, vq, ve in zip(vis, vis_qu, vis_es):
            if j != i and (fuzz.ratio(q, vq) >= umbral or fuzz.ratio(e, ve) >= 85):
                return True
        return False

    elegidas = set()
    por_sub = defaultdict(lambda: {"General": [], "Médico": []})
    for f in filas_b01:
        r = info.loc[f["ID"]]
        if not r["Quechua_borrador"]:
            continue
        if r["Dominio"] == "Médico" and int(r["Calidad"]) > CFG["control"]["calidad_max_medico"]:
            continue                                          # control médico solo con borradores de buena calidad
        if tiene_hermana(f["ID"]) or r["Código_original"] in excluir:
            continue
        por_sub[f["Sublote"]][r["Dominio"]].append(f["ID"])
    subs = sorted(por_sub)
    for dom, total in [("General", CFG["control"]["n_general"]), ("Médico", CFG["control"]["n_medico"])]:
        cupos = [total // len(subs) + (1 if k < total % len(subs) else 0) for k in range(len(subs))]
        for s, c in zip(subs, cupos):
            lista = sorted(por_sub[s][dom])
            elegidas |= set(rng.sample(lista, min(c, len(lista))))
    return elegidas


def elegir_test(filas: list[dict], C: pd.DataFrame, control: set[str], semilla: int, con_gemela: set[str]) -> set[str]:
    """P1: 10 % estratificado del bloque: 14 generales (proporcional a cada origen) + 6 médicas en unidades completas
    (pregunta + su respuesta), para no partir pares. No se sugiere una frase que tenga una casi gemela en train."""
    if not CFG["test_sugerido"]["activo"]:
        return set()
    rng = random.Random(semilla)
    info = C.set_index("ID")
    ids = [f["ID"] for f in filas if f["ID"] not in control]
    gen = [i for i in ids if info.loc[i, "Dominio"] == "General" and i not in con_gemela]
    out = set()
    por_pool = defaultdict(list)
    for i in gen:
        por_pool[info.loc[i, "Pool"]].append(i)
    n_gen = CFG["test_sugerido"]["n_general"]
    asignado = 0
    pools = sorted(por_pool, key=lambda p: -len(por_pool[p]))
    for k, p in enumerate(pools):
        n = round(n_gen * len(por_pool[p]) / len(gen)) if k < len(pools) - 1 else n_gen - asignado
        n = max(0, min(n, len(por_pool[p])))
        out |= set(rng.sample(sorted(por_pool[p]), n))
        asignado += n
    n_med = CFG["test_sugerido"]["n_medico"]
    unidades = defaultdict(list)
    for f in filas:
        if info.loc[f["ID"], "Dominio"] == "Médico":
            unidades[info.loc[f["ID"], "Plan_unidad"]].append(f["ID"])
    pares = sorted(u for u, lst in unidades.items()
                   if len(lst) == 2 and not (set(lst) & control) and not (set(lst) & con_gemela)
                   and {info.loc[i, "Rol"] for i in lst} == {"Médico", "Paciente"})
    for u in rng.sample(pares, min(n_med // 2, len(pares))):
        out |= set(unidades[u])
    return out


def ids_con_gemela(ids: list[str], C: pd.DataFrame) -> set[str]:
    """Filas que tienen otra casi igual (español) entre las filas planificadas: no sirven como prueba (P1)."""
    info = C.set_index("ID")
    claves = [clave_normalizada(info.loc[i, "Español"]) for i in ids]
    if not claves:
        return set()
    umbral = CFG["test_sugerido"]["similitud_max_train"]
    m = cdist(claves, claves, scorer=fuzz.ratio, score_cutoff=umbral, workers=-1) >= umbral
    np.fill_diagonal(m, False)
    return {ids[a] for a in np.where(m.any(axis=1))[0]}


# ======================================================================= salida: Word del traductor


def word_traductor(b: str, filas: list[dict], C: pd.DataFrame, QA: pd.DataFrame, control: set[str]):
    info = C.set_index("ID")
    avisos = defaultdict(list)
    for _, q in QA.iterrows():
        avisos[q["ID"]].append((q["Tipo"], q["Severidad"], q["Detalle"]))
    doc = W.documento_horizontal()
    W.pie_con_pagina(doc, f"Bloque {b}")
    h = doc.add_heading(f"Bloque {b} — Traducción al quechua chanka", level=1)
    doc.add_paragraph(f"Este documento tiene {len(filas)} oraciones en {CB['sublotes']} sublotes de {CB['tamano_sublote']}. "
                      "Gracias por su ayuda.")
    doc.add_heading("Qué hacer", level=2)
    for t in [
        "Lea cada oración en español y revise el quechua de la columna «Quechua». Corríjalo ahí mismo para que sea quechua chanka natural.",
        "Si la celda de quechua está vacía, escriba la traducción.",
        "Si una frase en español no se diría así, escriba cómo lo diría naturalmente un hablante y anótelo en «Observaciones».",
        "Algunas filas traen un aviso en gris en «Observaciones». Es solo una sospecha automática: usted decide. Escriba su comentario después del aviso.",
        "No borre, mueva ni cambie la columna ID. No agregue ni quite filas ni columnas.",
        "No use «Control de cambios». Si lo usó, antes de devolver el archivo vaya a Revisar → Aceptar todos los cambios.",
        "Guarde y devuelva este mismo archivo.",
    ]:
        doc.add_paragraph(t, style="List Number")
    doc.add_paragraph("«Quién habla»: Médico, Paciente, o — cuando es una oración general (sin hablante).")
    doc.add_heading("Convención ortográfica (alfabeto oficial del quechua chanka)", level=2)
    t = W.tabla(doc, ["Regla", "Escriba", "No escriba"], [9.0, 8.0, 8.0])
    for regla, si, no in [
        ("Tres vocales en palabras quechuas: a, i, u", "unquy, -yuq, qullqi", "onqoy, -yoq, qolqe"),
        ("Sin glotalizadas ni aspiradas", "pacha, qaway, achka", "p'acha, qhaway, achkha"),
        ("k, w, h en lugar de c, qu, hu, j", "wiksa, wasi, hamuy", "wicsa, huasi, jamuy"),
        ("Formas del chanka", "kachkanki, manam, ñuqanchik, -pti-", "kashanki, manan, ñoqanchis, -qti-"),
        ("Préstamos y nombres propios: como en español", "hospital, Ayacucho", "—"),
    ]:
        celdas = t.add_row().cells
        for c, v in zip(celdas, (regla, si, no)):
            W.texto_celda(c, v)
    doc.add_page_break()
    anchos = [1.2, 2.2, 2.3, 8.3, 8.6, 4.7]
    por_sub = defaultdict(list)
    for f in filas:
        por_sub[f["Sublote"]].append(f)
    for k, s in enumerate(sorted(por_sub)):
        fs = por_sub[s]
        if k > 0:
            doc.add_page_break()
        doc.add_heading(f"Sublote {k + 1} (filas {fs[0]['N°']}–{fs[-1]['N°']})", level=2)
        tb = W.tabla(doc, ["N°", "ID", "Quién habla", "Español", "Quechua (corrija o escriba aquí)", "Observaciones"], anchos)
        for f in fs:
            r = info.loc[f["ID"]]
            fila = tb.add_row()
            W.no_partir_fila(fila)
            c = fila.cells
            W.texto_celda(c[0], str(f["N°"]))
            W.texto_celda(c[1], f["ID"])
            W.texto_celda(c[2], r["Rol"] if r["Dominio"] == "Médico" else "—")
            W.texto_celda(c[3], r["Español"])
            qu = "" if f["ID"] in control else r["Quechua_borrador"]
            W.texto_celda(c[4], qu)
            av = "" if (f["ID"] in control or not r["Quechua_borrador"]) else aviso_word(avisos.get(f["ID"], []))
            W.texto_celda(c[5], av, cursiva=True, gris=True, tamano=11)
        W.anchos_columnas(tb, anchos)
    destino = BLOQUES_DIR / b / f"{b}_traductor.docx"
    asegurar_dir(destino.parent)
    guardar_seguro(destino, lambda tmp: doc.save(tmp))
    return destino


# ======================================================================= salida: Excel de trabajo


def excel_trabajo(b: str, filas: list[dict], C: pd.DataFrame, O: pd.DataFrame, QA: pd.DataFrame, control: set[str], test: set[str]):
    info = C.set_index("ID")
    org = O.set_index("ID")
    qa_txt = QA.groupby("ID").apply(lambda g: " | ".join(f"{t}/{s}: {d}" for t, s, d in zip(g["Tipo"], g["Severidad"], g["Detalle"])),
                                    include_groups=False).to_dict() if len(QA) else {}
    bloque_rows, detalle_rows = [], []
    for f in filas:
        r = info.loc[f["ID"]]
        o = org.loc[f["ID"]]
        qu = r["Quechua_validado"] or r["Quechua_borrador"]
        bloque_rows.append({"N°": f["N°"], "ID": f["ID"], "Sublote": f["Sublote"], "Dominio": r["Dominio"],
                            "Rol": r["Rol"] if r["Dominio"] == "Médico" else "—", "Español": r["Español"], "Quechua": qu,
                            "Estado": r["Estado"], "Grabado_ES": "☐", "Grabado_QU": "☐", "Notas": ""})
        detalle_rows.append({"N°": f["N°"], "ID": f["ID"], "Sublote": f["Sublote"], "Control": "Sí" if f["ID"] in control else "No",
                             "Test_sugerido": "Sí" if f["ID"] in test else "No", "Responde_a": r["Responde_a"],
                             "Quechua_borrador": r["Quechua_borrador"], "Quechua_validado": r["Quechua_validado"], "%_edición": "",
                             "Comentario_traductor": "", "Avisos_ortografía_traductor": "",
                             "Avisos_QA": qa_txt.get(f["ID"], ""), "Revisar_Juan": r["Revisar_Juan"], "Fuente": o["Fuente"],
                             "Código_original": o["Código_original"], "Tipo_borrador": o["Tipo_borrador"],
                             "Variedad_fuente": o["Variedad_fuente"], "Origen_verificado": o["Origen_verificado"],
                             "Estado_unificado": o["Estado_unificado"], "Observaciones_heredadas": o["Observaciones_heredadas"],
                             "Escenario": r["Escenario"], "Etapa_consulta": r["Etapa_consulta"]})
    B = pd.DataFrame(bloque_rows)
    D = pd.DataFrame(detalle_rows)
    destino = BLOQUES_DIR / b / f"{b}_trabajo.xlsx"
    conservar_trabajo(destino, B, D)
    resumen = resumen_bloque(b, filas, C, control, test)
    wb = Workbook()
    wb.remove(wb.active)
    escribir_hoja(wb, "Bloque", B)
    escribir_hoja(wb, "Detalle", D, estado_col=None, dominio_col=None)
    escribir_clave_valor(wb, "Resumen", resumen)
    wb.active = 0
    asegurar_dir(destino.parent)
    guardar_seguro(destino, lambda tmp: wb.save(tmp))
    return destino


COLS_EQUIPO_BLOQUE = ["Grabado_ES", "Grabado_QU", "Notas"]
COLS_EQUIPO_DETALLE = ["%_edición", "Comentario_traductor", "Avisos_ortografía_traductor"]


def conservar_trabajo(destino, B: pd.DataFrame, D: pd.DataFrame):
    """Si el Excel de trabajo ya existe, conserva lo que anotó el equipo o trajo el traductor (casillas de grabación, notas,
    % de edición, comentarios y avisos del traductor) para que regenerar no lo borre."""
    if not destino.exists():
        return
    try:
        viejo = pd.read_excel(destino, sheet_name=None, dtype=str, keep_default_na=False)
    except PermissionError:
        raise PermissionError(f"{destino} está abierto. Ciérralo y vuelve a correr con --forzar.")
    for hoja, df, cols in (("Bloque", B, COLS_EQUIPO_BLOQUE), ("Detalle", D, COLS_EQUIPO_DETALLE)):
        if hoja not in viejo or "ID" not in viejo[hoja]:
            continue
        v = viejo[hoja].drop_duplicates("ID").set_index("ID")
        for c in cols:
            if c in v:
                previo = df["ID"].map(v[c]).fillna("")
                vacio = "☐" if c.startswith("Grabado") else ""
                df[c] = [p if (p and p != vacio) else actual for p, actual in zip(previo, df[c])]


def resumen_bloque(b, filas, C, control, test) -> list[tuple]:
    info = C.set_index("ID")
    ids = [f["ID"] for f in filas]
    sub = defaultdict(lambda: Counter())
    for f in filas:
        r = info.loc[f["ID"]]
        sub[f["Sublote"]][r["Dominio"]] += 1
    dom = Counter(info.loc[ids, "Dominio"])
    med = info.loc[[i for i in ids if info.loc[i, "Dominio"] == "Médico"]]
    roles = Counter(med["Rol"])
    gen = info.loc[[i for i in ids if info.loc[i, "Dominio"] == "General"]]
    pools = Counter(gen["Pool"])
    sufijos = set()
    for q in info.loc[ids, "Quechua_borrador"]:
        sufijos |= Q.unidades_cobertura(q)["sufijo"]
    total_suf = len({s for v in Q.SUFIJOS_CATEGORIA.values() for s in v})
    temas = list(dict.fromkeys(t for t in med["Plan_tema"] if t))
    filas_out = [(f"Resumen del bloque {b}",), ("Generado", hoy()),
                 ("Total de oraciones", len(ids)),
                 ("Generales / médicas", f"{dom.get('General', 0)} / {dom.get('Médico', 0)} (objetivo {CB['generales']} / {CB['medicas']})"),
                 ("Médico / paciente (parte médica)", f"{roles.get('Médico', 0)} / {roles.get('Paciente', 0)}"),
                 ("Parte general por origen", ", ".join(f"{NOMBRE_POOL[p]}: {pools.get(p, 0)}" for p in POOLS + [POOL_CLAUDE])),
                 ("Tiempos y rasgos de la parte general", resumen_tiempos(list(gen.index), C)),
                 ("Etapas o temas de la parte médica", " → ".join(temas)),
                 ("Sufijos del quechua presentes (borradores)", f"{len(sufijos)} de {total_suf} (segmentación aproximada)"),
                 ("Oraciones de control (D7)", f"{len(control & set(ids))} (celda Quechua vacía en el Word)" if b == "B01" else "No aplica (solo B01)"),
                 ("Test sugerido (P1)", f"{len(test & set(ids))} filas marcadas; por ahora todo es train"),
                 ("Español redactado o adaptado por Claude (Revisar_Juan)", int((info.loc[ids, "Estado"] == "Revisar_Juan").sum())),
                 ("Sin borrador quechua (el traductor traduce desde cero)", int((info.loc[ids, "Quechua_borrador"] == "").sum())),
                 ("",), ("Por sublote",), ("Sublote", "Generales + médicas = total")]
    for s in sorted(sub):
        filas_out.append((s, f"{sub[s]['General']} + {sub[s]['Médico']} = {sub[s]['General'] + sub[s]['Médico']}"))
    return filas_out


# ======================================================================= guía de grabación (§8.6)


def guia_grabacion():
    a = CFG["audio"]
    doc = W.documento_vertical()
    W.pie_con_pagina(doc, "Guía de grabación")
    doc.add_heading("Guía de grabación — Corpus quechua chanka médico", level=1)
    doc.add_paragraph("Para el grabador y quien acompañe la sesión. Se graba con la hoja BXX_grabacion.docx, que se genera "
                      "cuando el traductor devuelve el bloque (script 07_hoja_grabacion.py).")
    doc.add_heading("1. Antes de empezar", level=2)
    for t in ["Sala silenciosa: puertas y ventanas cerradas, sin ventilador, radio ni televisor; celulares en silencio.",
              "Micrófono a unos 20 cm de la boca, siempre a la misma distancia. Haga una prueba de 10 segundos y escúchela.",
              f"Formato del audio: WAV, {a['frecuencia_hz'] // 1000} kHz, mono, {a['bits']} bits "
              "(en Audacity: Frecuencia del proyecto 16000 Hz, 1 canal; Exportar → WAV, PCM 16 bits con signo).",
              "Tenga a mano la hoja de grabación impresa y un lapicero para marcar y anotar."]:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_heading("2. Durante la grabación", level=2)
    for t in ["Grabe un sublote por vez (50 oraciones, cerca de 1 hora). Si la sesión se corta, se cierran sublotes completos.",
              "En cada fila lea primero el ESPAÑOL y después el QUECHUA, con una buena pausa (2 segundos) entre ambos.",
              "Deje 1 segundo de silencio al inicio y al final de cada audio. Hable con voz natural, sin apurarse.",
              "Si se equivoca, repita la oración completa (no solo la palabra).",
              "Marque ☐ ES y ☐ QU en la hoja cuando cada audio quede bien grabado."]:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_heading("3. Si se dice distinto a lo escrito (regla D8)", level=2)
    doc.add_paragraph("No se corrige al hablante. Si al grabar dice la frase de otra manera (otra palabra, otro orden), "
                      "anote en «Notas» exactamente lo que dijo. Después se transcribe lo que se dijo: la transcripción "
                      "tiene que coincidir con el audio. Solo es «gold» lo que el traductor validó y luego se grabó.")
    doc.add_heading("4. Nombre de los archivos", level=2)
    t = W.tabla(doc, ["Audio", "Nombre del archivo", "Ejemplo"], [4.0, 6.0, 7.0])
    for fila in [("Español", "{ID}_es.wav", "CLI-0011_es.wav"), ("Quechua", "{ID}_qu.wav", "CLI-0011_qu.wav")]:
        celdas = t.add_row().cells
        for c, v in zip(celdas, fila):
            W.texto_celda(c, v)
    doc.add_paragraph("")
    for t_ in ["Se nombran por ID y no por bloque, porque una fila puede cambiar de bloque antes de grabarse.",
               "Guarde todo en la carpeta 05_audio (puede usar una subcarpeta por sesión; el script busca por nombre).",
               "Un archivo por oración e idioma. Si graba seguido, diga el ID en voz alta antes de cada fila para poder cortar después."]:
        doc.add_paragraph(t_, style="List Bullet")
    doc.add_heading("5. Al terminar la sesión", level=2)
    for t_ in ["Copie los audios a 05_audio y una copia de respaldo en otro disco.",
               "Corra: .venv\\Scripts\\python.exe scripts\\08_manifest.py BXX — revisa que estén todos los audios, con el formato "
               "y la duración correctos, y arma 06_dataset/manifest_BXX.csv.",
               "Pase a limpio las notas de la hoja (frases dichas distinto) al Excel de trabajo del bloque."]:
        doc.add_paragraph(t_, style="List Bullet")
    doc.add_paragraph("Formato de audio (P5): confirmar con el equipo de modelo antes de la primera sesión.").runs[0].italic = True
    destino = BLOQUES_DIR / "GUIA_GRABACION.docx"
    asegurar_dir(destino.parent)
    guardar_seguro(destino, lambda tmp: doc.save(tmp))
    return destino


# ======================================================================= maestro


LEEME = [
    ("Corpus quechua chanka médico — archivo maestro",),
    ("Qué es", "La única fuente de verdad del corpus (INF657, PUCP). Dueño: Juan Zavala o quien él designe. Nadie trabaja sobre copias propias."),
    ("Generado", "Por los scripts de D:\\Corpus_Quechua\\scripts (cadena 00→05, f6). Fecha de los datos: {fecha}."),
    ("",),
    ("Pestañas",),
    ("Corpus", "Vista limpia: una fila por oración que se puede grabar."),
    ("Origen", "De dónde sale cada fila (fuente, archivo, hoja, código y fila original, tipo de borrador, variedad, verificación)."),
    ("Cambios", "Cada modificación hecha por los scripts: campo, antes, después, regla y fecha."),
    ("Descartes", "Todo lo que no entra, con su motivo y el ID que lo reemplaza (si es duplicado). Nada se borra sin dejar rastro aquí."),
    ("Lexico", "Palabras sueltas, con ID LEX-xxxx (no se graban solas). Las prioritarias tienen una frase portadora en el Corpus: "
               "la columna Frase_portadora tiene el ID (CLI-…) de esa frase."),
    ("QA", "Avisos de la revisión automática (ORTO, VARIEDAD, PRESTAMO, ALINEACION, SENTIDO, OOV, DUP, TRANSCRIPCION) con severidad "
           "ALTA/MEDIA/BAJA. Son sospechas, no correcciones. Al Word del traductor solo pasa 1 aviso ALTA por fila."),
    ("Bloques", "Una fila por bloque planificado: conteos (generales/médicas, médico/paciente, origen de la parte general, "
                "oraciones redactadas por Claude), mezcla de tiempos, reparto por sublote, temas médicos, control, test sugerido, "
                "filas para Revisar_Juan y estado (Planificado · Archivos generados · Congelado). La última fila («Sin asignar») "
                "dice cuánto material queda y si alcanza para otro bloque."),
    ("Traductor_metricas", "Una fila por bloque devuelto (scripts/06_importar_traductor.py): % sin cambios / editado / reescrito, "
                           "distancia al borrador, prueba de anclaje con las oraciones de control y la recomendación de seguir o "
                           "no mandando borradores."),
    ("",),
    ("Columnas de Corpus",),
    ("ID", "GEN-xxxx (general) o CLI-xxxx (médico/clínico); LEX-xxxx solo en Lexico. Un ID nunca cambia ni se reutiliza: el "
           "registro está en 02_maestro/ids.csv. Los bloques ya enviados quedan congelados (no se rearman)."),
    ("Dominio", "General o Médico (70 % / 30 % de cada bloque)."),
    ("Escenario / Etapa_consulta", "Sección y subsección de origen / etapa de la consulta (Calgary-Cambridge, §11.3)."),
    ("Rol", "Médico / Paciente / — (sin hablante)."),
    ("Quechua_borrador", "Borrador NO validado (traductor automático, normalizado de otra variedad, libro…). Nunca es gold."),
    ("Quechua_validado", "Lo que devolvió el traductor. Gold = validado y grabado (D8)."),
    ("Estado", "Borrador · Sin borrador · Revisar_Juan (español redactado o adaptado por Claude) · Enviado · Validado · Grabado."),
    ("Prioridad", "Número del bloque planificado (1 = B01 … 5 = B05); vacío = aún sin bloque."),
    ("Bloque / Sublote / N_bloque", "B01…; B01-S1…B01-S4; posición 1–200 dentro del bloque."),
    ("Split", "train (todo es entrenamiento, P1)."),
    ("Test_sugerido", "Sí = parte del 10 % estratificado que se podría apartar como prueba si Juan lo confirma (P1)."),
    ("Control", "Sí = oración de control de B01 (D7): va al traductor con la celda Quechua vacía."),
    ("Responde_a", "En respuestas del paciente: ID de la pregunta del médico (van en el mismo sublote)."),
    ("",),
    ("Reglas clave",),
    ("D8", "Ningún borrador entra al entrenamiento ni a la evaluación como verdad."),
    ("Ortografía", "Alfabeto oficial del quechua chanka (§10.3). Los scripts solo hicieron cambios mecánicos (registrados en Cambios)."),
]


def escribir_maestro(C, O, CBios, D, L, QA, BL, TM):
    wb = Workbook()
    wb.remove(wb.active)
    escribir_clave_valor(wb, "LEEME", [tuple(x.format(fecha=hoy()) if isinstance(x, str) else x for x in f) for f in LEEME])
    escribir_hoja(wb, "Corpus", C[COLS_CORPUS])
    escribir_hoja(wb, "Origen", O, estado_col=None, dominio_col=None)
    escribir_hoja(wb, "Cambios", CBios, estado_col=None, dominio_col=None)
    escribir_hoja(wb, "Descartes", D, estado_col=None, dominio_col=None)
    escribir_hoja(wb, "Lexico", L, estado_col=None, dominio_col=None)
    escribir_hoja(wb, "QA", QA, estado_col=None, dominio_col=None)
    escribir_hoja(wb, "Bloques", BL, estado_col=None, dominio_col=None)
    escribir_hoja(wb, "Traductor_metricas", TM, estado_col=None, dominio_col=None)
    wb.active = 1
    asegurar_dir(MAESTRO.parent)
    guardar_seguro(MAESTRO, lambda tmp: wb.save(tmp))


COLS_TM = ["Bloque", "Fecha_importación", "Filas", "Filas_con_borrador", "Sin_cambios_%", "Editado_%", "Reescrito_%",
           "Distancia_media_borrador", "Control_n", "Distancia_media_control", "Anclaje", "Recomendación", "Faltantes", "Sobrantes",
           "Duplicados"]


ESTADOS_ENVIADOS = ("Enviado", "Validado", "Grabado")


def leer_maestro_viejo() -> dict[str, pd.DataFrame]:
    if not MAESTRO.exists():
        return {}
    try:
        return pd.read_excel(MAESTRO, sheet_name=None, dtype=str, keep_default_na=False)
    except PermissionError:
        raise PermissionError(f"{MAESTRO} está abierto. Ciérralo y vuelve a correr.")


def bloques_congelados(viejo: dict[str, pd.DataFrame]) -> dict[str, dict]:
    """H5: un bloque que ya salió al traductor no se vuelve a armar: se copia su composición (IDs, N°, sublote, control y
    test) del maestro anterior. Se congela si está en config.bloque.congelados o si alguna de sus filas ya está Enviado,
    Validado o Grabado."""
    cong = set(CB.get("congelados") or [])
    V = viejo.get("Corpus")
    if V is None or "Bloque" not in V:
        if cong:
            raise ValueError(f"config.bloque.congelados = {sorted(cong)}, pero no hay maestro anterior del que copiar su composición.")
        return {}
    cong |= set(V.loc[V["Estado"].isin(ESTADOS_ENVIADOS) & (V["Bloque"] != ""), "Bloque"])
    salida = {}
    for b in sorted(cong):
        S = V[V["Bloque"] == b].copy()
        if S.empty:
            raise ValueError(f"El bloque {b} está congelado pero no aparece en el maestro anterior.")
        S["_n"] = S["N_bloque"].astype(int)
        S = S.sort_values("_n")
        salida[b] = {"filas": [{"ID": i, "N°": int(n), "Sublote": s} for i, n, s in zip(S["ID"], S["_n"], S["Sublote"])],
                     "control": set(S.loc[S["Control"] == "Sí", "ID"]), "test": set(S.loc[S["Test_sugerido"] == "Sí", "ID"])}
    return salida


def conservar_validados(C: pd.DataFrame, viejo: dict[str, pd.DataFrame]) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Si el maestro ya existe, conserva Quechua_validado, Estado (Enviado, Validado, Grabado) y Traductor_metricas.
    H1: el quechua validado solo se conserva si el español de la fila sigue igual; si el español cambió después de enviarlo,
    la fila vuelve a Revisar_Juan con una nota (hay que reenviarla al traductor)."""
    TM = pd.DataFrame(columns=COLS_TM)
    if "Corpus" in viejo:
        v = viejo["Corpus"].drop_duplicates("ID").set_index("ID")
        col = lambda campo: C["ID"].map(v[campo]).fillna("") if campo in v else pd.Series([""] * len(C), index=C.index)
        es_v, qu_v, est_v = col("Español"), col("Quechua_validado"), col("Estado")
        qv, est, nota = list(C["Quechua_validado"]), list(C["Estado"]), list(C["Revisar_Juan"])
        cam = viejo.get("Cambios", pd.DataFrame(columns=["ID", "Campo"]))
        propuesta = set(cam.loc[cam["Campo"] == "Español (propuesta del traductor)", "ID"])
        for k, (i, es, e0, q0, s0) in enumerate(zip(C["ID"], C["Español"], es_v, qu_v, est_v)):
            if s0 == "Revisar_Juan" and i in propuesta and e0 == es:
                est[k] = "Revisar_Juan"            # el traductor propuso otro español (06) y Juan aún no decide
                nota[k] = "El traductor propuso otro español al devolver el bloque: ver la pestaña Cambios."
                continue
            if s0 not in ESTADOS_ENVIADOS:
                continue
            if e0 == es:
                qv[k] = qv[k] or q0
                est[k] = s0
            else:
                est[k] = "Revisar_Juan"
                nota[k] = (f"El español cambió después de enviarlo al traductor (antes: «{e0}»"
                           + (f"; quechua validado de entonces: «{q0}»" if q0 else "") + "): reenviar esta fila.")
        C["Quechua_validado"], C["Estado"], C["Revisar_Juan"] = qv, est, nota
    if "Traductor_metricas" in viejo and len(viejo["Traductor_metricas"]):
        TM = viejo["Traductor_metricas"]
    return C, TM


# ======================================================================= lista para Juan


DONDE_F5 = {"B01": "`plan_b01.yaml`", "resp": "`respuestas_paciente.yaml`", "port": "`frases_portadoras.yaml`",
            "gen": "`oraciones_generales.yaml`"}


def celda_md(s: str) -> str:
    return str(s).replace("|", "\\|").replace("\n", " ")


def revision_juan(C: pd.DataFrame, cam: pd.DataFrame, filas_bloque: dict[str, list[dict]]):
    """04_reportes/REVISION_JUAN.md: filas con Estado = Revisar_Juan de los bloques con archivos, con el archivo de
    scripts/datos/ donde se corrige cada una. En las adaptadas o partidas se muestra el borrador quechua, para ver si
    sigue diciendo lo mismo que el español nuevo (H17)."""
    info = C.set_index("ID")
    md = [f"# Revisión de Juan — español redactado o adaptado por Claude ({', '.join(CB['bloques_con_archivos'])})", "",
          "Estas filas tienen Estado = **Revisar_Juan**. Revísalas antes de enviar el Word al traductor. Si cambias algo, "
          "edita el archivo de `scripts/datos/` indicado y vuelve a correr `04_limpieza` → `05_huecos` → `f6_armar_bloques` "
          "(con `--forzar`). Las oraciones nuevas van al traductor SOLO en español: él escribe el quechua.", ""]
    for b in CB["bloques_con_archivos"]:
        ids = [f["ID"] for f in filas_bloque.get(b, [])]
        rj = info.loc[ids]
        rj = rj[rj["Estado"] == "Revisar_Juan"]
        md += [f"## {b}: {len(rj)} filas", "",
               "| N° | ID | Rol | Español (actual) | Original / motivo | Borrador quechua | Dónde se edita |",
               "|---|---|---|---|---|---|---|"]
        for i, r in rj.iterrows():
            nota = r["Revisar_Juan"]
            if nota.startswith("El español cambió"):
                orig, donde = nota, "reenviar al traductor"
            elif r["uid"].startswith("F5:"):
                tipo = r["uid"].split(":")[1]
                orig = "— (oración general nueva, varios tiempos)" if tipo == "gen" else "— (frase nueva de F5)"
                donde = DONDE_F5.get(tipo, "`scripts/datos/`")
            else:
                ccs = cam[cam["ID"] == i]
                c0 = ccs[ccs["Regla"].str.contains("adaptación|partir oración", regex=True)]
                orig = (c0["Antes"].iloc[0] + " · " + c0["Regla"].iloc[0]) if len(c0) else nota or "—"
                donde = "`adaptaciones_es.yaml`" if "adaptación" in orig else "`particiones_es.yaml`"
            qu = r["Quechua_borrador"] if not r["uid"].startswith("F5:") else ""
            md.append("| " + " | ".join(celda_md(x) for x in (r["N_bloque"], i, r["Rol"], r["Español"], orig, qu or "—", donde)) + " |")
        md.append("")
    escribir_texto(ruta("reportes") / "REVISION_JUAN.md", "\n".join(md))


# ======================================================================= principal


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--forzar", action="store_true")
    ap.add_argument("--solo-archivos", default="", help="regenerar solo los archivos de estos bloques (p. ej. B03,B04); "
                    "útil si otro archivo está abierto en Word/LibreOffice")
    a = ap.parse_args()
    solo = [x.strip().upper() for x in a.solo_archivos.split(",") if x.strip()]
    if CK.exists() and MAESTRO.exists() and not a.forzar:
        print(f"Ya existen {CK.relative_to(RAIZ)} y el maestro; usa --forzar.")
        return 0

    H = cargar_checkpoint("F5_huecos")
    C = H["Corpus"].copy()
    O = H["Origen"].copy()
    QA = H["QA"].copy()
    viejo = leer_maestro_viejo()
    cong = bloques_congelados(viejo)
    ids_cong = {f["ID"] for b in cong for f in cong[b]["filas"]}
    faltan = sorted(ids_cong - set(C["ID"]))
    if faltan:
        raise ValueError(f"Filas de bloques ya enviados ({', '.join(sorted(cong))}) ya no están en el corpus: {', '.join(faltan[:10])}"
                         f"{' …' if len(faltan) > 10 else ''}. Su texto cambió o se descartó: restaurarlo o tratarlo a mano.")
    bloques_plan = sorted(b for b in set(C["Plan_bloque"]) | set(cong) if b)
    libres = [b for b in bloques_plan if b not in cong]
    gen_sel, notas = seleccionar_general(C, libres, excluir=ids_cong)
    bloques = [b for b in bloques_plan if b in gen_sel or b in cong]

    filas_bloque = {b: cong[b]["filas"] if b in cong else armar_bloque(b, gen_sel[b], C) for b in bloques}
    if "B01" in cong:
        control = set(cong["B01"]["control"])
    elif "B01" in filas_bloque:
        # B01 y B02 se entregan juntos: el control no puede tener una «hermana» con borrador en ninguno de los dos
        visibles = [f["ID"] for b in ("B01", "B02") for f in filas_bloque.get(b, [])]
        control = elegir_control(filas_bloque["B01"], C, visibles)
    else:
        control = set()
    planificadas = [f["ID"] for b in bloques for f in filas_bloque[b]]
    con_gemela = ids_con_gemela(planificadas, C)
    test = set()
    for k, b in enumerate(bloques):
        test |= cong[b]["test"] if b in cong else elegir_test(filas_bloque[b], C, control, CFG["semilla"] + k, con_gemela)

    # columnas de plan en el corpus
    asign = {f["ID"]: (b, f["Sublote"], f["N°"]) for b in bloques for f in filas_bloque[b]}
    C["Bloque"] = C["ID"].map(lambda i: asign.get(i, ("", "", ""))[0])
    C["Sublote"] = C["ID"].map(lambda i: asign.get(i, ("", "", ""))[1])
    C["N_bloque"] = C["ID"].map(lambda i: asign.get(i, ("", "", ""))[2])
    C["Prioridad"] = C["Bloque"].map(lambda b: str(int(b[1:])) if b else "")
    C["Control"] = C["ID"].map(lambda i: "Sí" if i in control else "No")
    C["Test_sugerido"] = C["ID"].map(lambda i: "Sí" if i in test else "No")
    C["Split"] = "train"
    C, TM = conservar_validados(C, viejo)
    orden_dom = {"General": 0, "Médico": 1}
    C = C.sort_values(by=["Dominio", "ID"], key=lambda s: s.map(orden_dom) if s.name == "Dominio" else s).reset_index(drop=True)

    # pestaña Bloques
    filas_bl = []
    for b in bloques:
        ids = [f["ID"] for f in filas_bloque[b]]
        sub = C.set_index("ID").loc[ids]
        med = sub[sub["Dominio"] == "Médico"]
        gen = sub[sub["Dominio"] == "General"]
        cnt_sub = Counter(zip(sub["Sublote"], sub["Dominio"]))
        filas_bl.append({
            "Bloque": b, "Total": len(ids), "Generales": len(gen), "Médicas": len(med),
            "Médico": int((med["Rol"] == "Médico").sum()), "Paciente": int((med["Rol"] == "Paciente").sum()),
            "Principito": int((gen["Pool"] == "principito").sum()), "G2_frases_variadas": int((gen["Pool"] == "g2_diccionario").sum()),
            "G2_ejercicios": int((gen["Pool"] == "g2_ejercicio").sum()),
            "Generales_Claude": int((gen["Pool"] == POOL_CLAUDE).sum()),
            "Médicas_Claude": int(med["uid"].str.startswith("F5:").sum()),
            "Tiempos_general": resumen_tiempos(list(gen.index), C),
            "Sublotes (gen+méd)": " | ".join(f"S{k}: {cnt_sub[(f'{b}-S{k}', 'General')]}+{cnt_sub[(f'{b}-S{k}', 'Médico')]}" for k in range(1, CB["sublotes"] + 1)),
            "Etapas/temas": " → ".join(dict.fromkeys(t for t in med["Plan_tema"] if t)),
            "Control": int((sub["Control"] == "Sí").sum()), "Test_sugerido": int((sub["Test_sugerido"] == "Sí").sum()),
            "Revisar_Juan": int((sub["Estado"] == "Revisar_Juan").sum()),
            "Sin_borrador": int((sub["Quechua_borrador"] == "").sum()),
            "Estado": ("Congelado (ya salió al traductor; composición fija)" if b in cong else
                       "Archivos generados (Word del traductor y Excel de trabajo)" if b in CB["bloques_con_archivos"] else "Planificado"),
            "Notas": " ".join(n for n in notas if n.startswith(b)),
        })
    resto = C[C["Bloque"] == ""]
    r_gen, r_med = resto[resto["Dominio"] == "General"], resto[resto["Dominio"] == "Médico"]
    por_pool = Counter(r_gen["Pool"])
    usables = sum(por_pool.get(p, 0) for p in ("principito", "g2_diccionario", POOL_CLAUDE)) + \
        min(por_pool.get("g2_ejercicio", 0), CFG["general"]["cuotas_por_bloque"]["g2_ejercicio_max"])
    nota_resto = (f"Quedan {len(r_gen)} generales ({', '.join(f'{NOMBRE_POOL[p]}: {por_pool.get(p, 0)}' for p in POOLS + [POOL_CLAUDE])}) "
                  f"y {len(r_med)} médicas sin bloque. ")
    nota_resto += (f"Alcanza para otro bloque de {CB['generales']} generales respetando el tope de ejercicios."
                   if usables >= CB["generales"] else
                   f"Respetando el tope de ejercicios solo hay {usables} generales usables: no alcanza para otro bloque de "
                   f"{CB['generales']}. Para seguir: más oraciones redactadas (oraciones_generales.yaml) o decidir P4 (AmericasNLP).")
    filas_bl.append({"Bloque": "Sin asignar", "Total": len(resto), "Generales": len(r_gen), "Médicas": len(r_med),
                     "Estado": "Material disponible para bloques futuros", "Notas": nota_resto})
    BL = pd.DataFrame(filas_bl).astype(object).fillna("")
    BL = BL.map(lambda v: int(v) if isinstance(v, float) and v.is_integer() else v)       # conteos enteros (no 32.0)

    # Origen solo para filas del corpus; Lexico y demás
    O = O[O["ID"].isin(C["ID"])].copy()
    O = O.set_index("ID").loc[C["ID"]].reset_index()[COLS_ORIGEN]
    CBios = H["Cambios"].drop(columns=["uid"], errors="ignore")
    if "Cambios" in viejo and "Regla" in viejo["Cambios"]:
        # lo que registraron 06 (importación, envío) no sale de la cadena 00→05: se conserva del maestro anterior
        del_06 = viejo["Cambios"][viejo["Cambios"]["Regla"].str.contains("(06", regex=False)]
        CBios = pd.concat([CBios, del_06[[c for c in CBios.columns if c in del_06.columns]]], ignore_index=True).fillna("")
    D = H["Descartes"].drop(columns=["uid", "Reemplazado_por_uid"], errors="ignore")
    L = H["Lexico"].drop(columns=["uid"], errors="ignore")

    escribir_maestro(C, O, CBios, D, L, QA, BL, TM)

    for b in CB["bloques_con_archivos"]:
        if b not in filas_bloque or (solo and b not in solo):
            continue
        if b in cong:
            print(f"{b}: congelado (ya salió al traductor): no se regeneran su Word ni su Excel de trabajo.")
            continue
        excel_trabajo(b, filas_bloque[b], C, O, QA, control, test)
        word_traductor(b, filas_bloque[b], C, QA, control)
    guia_grabacion()
    revision_juan(C, H["Cambios"], filas_bloque)

    # El checkpoint se escribe AL FINAL: es la marca de corrida completa (H8). Si algo falla antes (p. ej. un Word abierto),
    # no existe y la siguiente corrida vuelve a hacerlo todo aunque no se pase --forzar.
    hojas_ck = {"Corpus": C, "Origen": O, "Cambios": CBios, "Descartes": D, "Lexico": L, "QA": QA, "Bloques": BL,
                "Traductor_metricas": TM,
                "Filas_bloque": pd.DataFrame([dict(f, Bloque=b) for b in bloques for f in filas_bloque[b]])}
    guardar_excel(CK, hojas_ck)

    print(f"OK → maestro, {CK.relative_to(RAIZ)}, archivos de {', '.join(CB['bloques_con_archivos'])}")
    print(BL[["Bloque", "Total", "Generales", "Médicas", "Médico", "Paciente", "Principito", "G2_frases_variadas", "G2_ejercicios",
              "Generales_Claude", "Médicas_Claude", "Sublotes (gen+méd)", "Control", "Test_sugerido", "Revisar_Juan"]].to_string(index=False))
    for _, r in BL.iterrows():
        if r["Tiempos_general"]:
            print(f"  {r['Bloque']}: {r['Tiempos_general']}")
    for n in notas:
        print("NOTA:", n)
    return 0


if __name__ == "__main__":
    sys.exit(main())
