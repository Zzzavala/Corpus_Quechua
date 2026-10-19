"""Crea el piloto humano de B01-S1 sin modificar el maestro.

Genera dos libros:
- REVISOR: se puede enviar; oculta los borradores de control y la procedencia.
- COORDINACION_NO_ENVIAR: contiene trazabilidad y borradores ocultos para análisis interno.
"""
from pathlib import Path
from collections import Counter
import hashlib
import json
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'04_reportes/eda_postseleccion/piloto_validacion_50'
MASTER=ROOT/'02_maestro/corpus_maestro.xlsx'
BLOQUE='B01';SUBLOTE='B01-S1'
AZUL='D9EAF7';AMARILLO='FFF2CC';GRIS='E7E6E6';VERDE='E2F0D9';ROJO='F4CCCC';BLANCO='FFFFFF'

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def texto(v):return '' if pd.isna(v) else str(v).strip()

def setup(ws,headers,widths):
    ws.append(headers);ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}1'
    thin=Side(style='thin',color='B7B7B7')
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor=AZUL);c.font=Font(bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
        c.border=Border(bottom=thin)
    ws.row_dimensions[1].height=32
    for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
    ws.sheet_view.showGridLines=False

def add_validations(ws,n):
    judgement=DataValidation(type='list',formula1='"Sí,No,Dudoso,No aplica"',allow_blank=True)
    decision=DataValidation(type='list',formula1='"Aprobar,Corregir español,Corregir quechua,Retraducir,Posponer"',allow_blank=True)
    ws.add_data_validation(judgement);ws.add_data_validation(decision)
    for col in ['H','I','J','K','L','M','N','O','P']:
        judgement.add(f'{col}2:{col}{n+1}')
    decision.add(f'Q2:Q{n+1}')

def instrucciones(wb,filas):
    ws=wb.active;ws.title='INSTRUCCIONES';ws.sheet_view.showGridLines=False
    contenido=[
        ('Piloto de validación humana — B01-S1','Título'),
        (f'{len(filas)} frases: {int(filas.Dominio.eq("General").sum())} generales y {int(filas.Dominio.ne("General").sum())} médicas.','Normal'),
        ('Objetivo','Sección'),
        ('Revisar primero el español y luego la equivalencia/naturalidad del quechua chanka. Este archivo no valida automáticamente ninguna fila.','Normal'),
        ('Cómo trabajar','Sección'),
        ('1. No cambie el ID ni el español directamente. Si necesita cambiar el español, seleccione «Corregir español» y escriba la propuesta en Comentario.','Normal'),
        ('2. En Quechua_trabajo corrija el borrador o escriba una traducción desde cero. Una celda vacía puede ser una frase de control o una frase sin borrador.','Normal'),
        ('3. Evalúe cada criterio con Sí, No, Dudoso o No aplica. «Sí» significa que el criterio se cumple.','Normal'),
        ('4. En Decisión use Aprobar, Corregir español, Corregir quechua, Retraducir o Posponer. Complete Revisor y Fecha.','Normal'),
        ('5. En frases médicas, Adecuación_clínica requiere una revisión competente. Si no corresponde a su rol, seleccione No aplica.','Normal'),
        ('6. No grabe todavía. La coordinación conciliará este libro con el maestro y resolverá desacuerdos.','Normal'),
        ('Regla de control','Sección'),
        ('Algunas celdas quechuas están vacías deliberadamente para obtener traducción libre. No solicite ni consulte otro borrador para esas filas durante esta revisión.','Advertencia'),
        ('Significado de criterios','Sección'),
        ('Claridad_ES: se entiende la intención. Naturalidad_ES: una persona diría la frase en la situación indicada. Autónomo_sin_contexto: se entiende sin una oración previa. Fácil_de_leer: puede pronunciarse sin dificultad innecesaria. Utilidad: aporta al corpus. Equivalencia_ES_QU: conserva el significado. Chanka_natural: la forma es natural en la variedad acordada. Ortografía_acordada: respeta la convención del proyecto.','Normal'),
    ]
    for i,(value,kind) in enumerate(contenido,1):
        ws.cell(i,1,value);ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=6)
        cell=ws.cell(i,1);cell.alignment=Alignment(wrap_text=True,vertical='top')
        if kind=='Título':cell.font=Font(size=16,bold=True,color='1F4E78');cell.fill=PatternFill('solid',fgColor=AZUL);ws.row_dimensions[i].height=28
        elif kind=='Sección':cell.font=Font(size=12,bold=True,color='1F4E78');cell.fill=PatternFill('solid',fgColor='DDEBF7')
        elif kind=='Advertencia':cell.font=Font(bold=True,color='9C0006');cell.fill=PatternFill('solid',fgColor=ROJO);ws.row_dimensions[i].height=42
        else:ws.row_dimensions[i].height=34
    ws.column_dimensions['A'].width=120

def crear_revisor(filas,path):
    wb=Workbook();instrucciones(wb,filas);ws=wb.create_sheet('REVISION')
    headers=['N','ID','Dominio','Quién_habla','Español','Modo_revisión','Quechua_trabajo',
             'Claridad_ES','Naturalidad_ES','Autónomo_sin_contexto','Fácil_de_leer','Utilidad',
             'Equivalencia_ES_QU','Chanka_natural','Ortografía_acordada','Adecuación_clínica',
             'Decisión','Comentario','Revisor','Fecha']
    setup(ws,headers,[6,13,11,13,48,22,48,13,15,20,15,12,18,15,18,19,20,42,20,14])
    for _,r in filas.iterrows():
        control=r.Control=='Sí';draft=texto(r.Quechua_borrador)
        visible='' if control else draft
        mode='Traducción libre (control)' if control else ('Revisar borrador' if draft else 'Traducir desde cero')
        ws.append([int(r.N_bloque),r.ID,r.Dominio,r.Rol,r.Español,mode,visible]+['']*13)
    add_validations(ws,len(filas))
    for row in ws.iter_rows(min_row=2,max_row=len(filas)+1):
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        for c in row[:6]:c.fill=PatternFill('solid',fgColor=GRIS)
        for c in row[6:]:c.fill=PatternFill('solid',fgColor=AMARILLO)
        row[0].alignment=Alignment(horizontal='center',vertical='top')
        ws.row_dimensions[row[0].row].height=54
    ws.auto_filter.ref=f'A1:T{len(filas)+1}'
    cod=wb.create_sheet('CÓDIGOS')
    setup(cod,['Campo','Valores','Interpretación'],[28,65,85])
    codes=[
        ('Criterios','Sí | No | Dudoso | No aplica','Sí = cumple el criterio; no significa aprobación total.'),
        ('Decisión','Aprobar','Español y quechua pueden pasar a conciliación final.'),
        ('Decisión','Corregir español','Escribir propuesta y motivo en Comentario; la traducción debe revisarse de nuevo.'),
        ('Decisión','Corregir quechua','Editar Quechua_trabajo y explicar el cambio relevante.'),
        ('Decisión','Retraducir','El par necesita una nueva traducción desde el español aprobado.'),
        ('Decisión','Posponer','Falta contexto o revisión especializada.'),
    ]
    for r in codes:cod.append(r)
    for row in cod.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
    wb.save(path)

def crear_coordinacion(filas,qa,path):
    wb=Workbook();ws=wb.active;ws.title='SEGUIMIENTO_INTERNO_NO_ENVIAR'
    headers=['N','ID','Dominio','Rol','Español','Control','Modo_revisor','Quechua_borrador_oculto',
             'Estado','Fuente','Tipo_borrador','Origen_verificado','Estado_unificado','Responde_a',
             'QA_tipos','QA_detalle','Decisión_revisor','Quechua_devuelto','Comentario_revisor','Conciliación','Responsable','Fecha']
    setup(ws,headers,[6,13,11,12,48,10,22,48,14,35,30,18,28,14,18,55,20,48,45,28,18,14])
    for _,r in filas.iterrows():
        control=r.Control=='Sí';draft=texto(r.Quechua_borrador)
        mode='Traducción libre (control)' if control else ('Revisar borrador' if draft else 'Traducir desde cero')
        av=qa.get(r.ID,{'tipos':'','detalle':''})
        ws.append([int(r.N_bloque),r.ID,r.Dominio,r.Rol,r.Español,r.Control,mode,draft,r.Estado,r.Fuente,
                   r.Tipo_borrador,r.Origen_verificado,r.Estado_unificado,r.Responde_a,av['tipos'],av['detalle']]+['']*6)
    for row in ws.iter_rows(min_row=2,max_row=len(filas)+1):
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        for c in row[:16]:c.fill=PatternFill('solid',fgColor=GRIS)
        for c in row[16:]:c.fill=PatternFill('solid',fgColor=AMARILLO)
        ws.row_dimensions[row[0].row].height=56
    ws.auto_filter.ref=f'A1:V{len(filas)+1}'
    res=wb.create_sheet('RESUMEN')
    setup(res,['Indicador','Valor'],[36,70])
    summary=[('Bloque',BLOQUE),('Sublote',SUBLOTE),('Filas',len(filas)),
        ('General',int(filas.Dominio.eq('General').sum())),('Médico',int(filas.Dominio.ne('General').sum())),
        ('Controles ciegos',int(filas.Control.eq('Sí').sum())),('Sin borrador original',int(filas.Quechua_borrador.eq('').sum())),
        ('Con borrador original',int(filas.Quechua_borrador.ne('').sum())),('Test_sugerido',int(filas.Test_sugerido.eq('Sí').sum())),
        ('Quechua validado en maestro',int(filas.Quechua_validado.ne('').sum())),('Estado','Piloto preparado; revisión humana pendiente')]
    for row in summary:res.append(row)
    note=wb.create_sheet('ADVERTENCIA')
    note['A1']='NO ENVIAR ESTE LIBRO AL TRADUCTOR/REVISOR'
    note['A2']='Contiene los borradores ocultos de las frases de control, procedencia y alertas internas. Use PILOTO_B01_S1_REVISOR.xlsx para la revisión ciega.'
    note['A1'].font=Font(size=16,bold=True,color='9C0006');note['A1'].fill=PatternFill('solid',fgColor=ROJO)
    note['A2'].alignment=Alignment(wrap_text=True);note.column_dimensions['A'].width=120;note.row_dimensions[2].height=50
    wb.save(path)

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    before=sha(MASTER)
    h={k:v.fillna('') for k,v in pd.read_excel(MASTER,sheet_name=None,dtype=str).items()}
    corpus=h['Corpus'].merge(h['Origen'],on='ID',validate='one_to_one')
    filas=corpus[(corpus.Bloque.eq(BLOQUE))&(corpus.Sublote.eq(SUBLOTE))].copy()
    filas['N_num']=pd.to_numeric(filas.N_bloque,errors='raise');filas=filas.sort_values('N_num')
    assert len(filas)==50 and filas.ID.is_unique
    assert Counter(filas.Dominio)==Counter({'General':34,'Médico':16})
    assert Counter(filas.Rol)==Counter({'—':34,'Médico':8,'Paciente':8})
    assert int(filas.Control.eq('Sí').sum())==5
    qa={}
    for cid,g in h['QA'][h['QA'].ID.isin(filas.ID)].groupby('ID'):
        qa[cid]={'tipos':' | '.join(sorted(set(g.Tipo))),
                 'detalle':' | '.join(f'{r.Severidad}: {r.Detalle}' for r in g.itertuples())}
    reviewer=OUT/'PILOTO_B01_S1_REVISOR.xlsx'
    internal=OUT/'PILOTO_B01_S1_COORDINACION_NO_ENVIAR.xlsx'
    crear_revisor(filas,reviewer);crear_coordinacion(filas,qa,internal)
    # Verificación de que ningún borrador de control está en el libro enviable.
    rev=pd.read_excel(reviewer,sheet_name='REVISION',dtype=str).fillna('')
    controls=set(filas[filas.Control.eq('Sí')].ID)
    sin_borrador=set(filas[filas.Quechua_borrador.eq('')].ID)
    assert len(rev)==50 and set(rev.ID)==set(filas.ID)
    assert rev[rev.ID.isin(controls)].Quechua_trabajo.eq('').all()
    assert rev[rev.ID.isin(controls)].Modo_revisión.eq('Traducción libre (control)').all()
    assert not (controls & sin_borrador), 'Una fila de control debe tener borrador oculto para poder comparar.'
    assert set(rev[rev.Quechua_trabajo.eq('')].ID)==controls|sin_borrador
    assert sha(MASTER)==before
    manifest={'maestro_sha256':before,'bloque':BLOQUE,'sublote':SUBLOTE,'filas':len(filas),
        'general':34,'medico':16,'medico_roles':{'Médico':8,'Paciente':8},'controles_ciegos':5,
        'sin_borrador_original':len(sin_borrador),'celdas_quechua_vacias_en_revisor':len(controls|sin_borrador),
        'desglose_vacias_revisor':'5 controles con borrador oculto + 14 sin borrador original',
        'quechua_validado':int(filas.Quechua_validado.ne('').sum()),
        'archivo_revisor_sha256':sha(reviewer),'archivo_interno_sha256':sha(internal),
        'borradores_control_expuestos_en_revisor':False,'maestro_modificado':False,
        'estado':'Preparado; pendiente de revisión humana y conciliación'}
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8');main()
