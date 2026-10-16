"""Formato de los Excel (§8.3): encabezado en negrita con color suave, primera fila inmovilizada, autofiltro,
anchos razonables, ajuste de texto, lista desplegable en Estado y color suave por Dominio. Solo valores."""
from __future__ import annotations

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ENCABEZADO = PatternFill("solid", fgColor="DCE8F7")
COLOR_DOMINIO = {"General": PatternFill("solid", fgColor="F4F7FB"), "Médico": PatternFill("solid", fgColor="FDF4EC")}
LINEA = Side(style="thin", color="D9D9D9")
BORDE = Border(bottom=LINEA)
ESTADOS = ["Borrador", "Sin borrador", "Revisar_Juan", "Enviado", "Validado", "Grabado"]
ANCHOS = {"N°": 6, "ID": 11, "Sublote": 9, "Dominio": 10, "Rol": 10, "Español": 55, "Quechua": 55, "Estado": 14,
          "Grabado_ES": 13, "Grabado_QU": 13, "Notas": 30, "Comentario_traductor": 40, "Avisos_ortografía_traductor": 40, "Quechua_borrador": 50, "Quechua_validado": 50,
          "Escenario": 40, "Etapa_consulta": 24, "Prioridad": 9, "Bloque": 8, "Split": 7, "Test_sugerido": 12,
          "Observaciones_heredadas": 50, "Detalle": 70, "Antes": 50, "Después": 50, "Regla": 40, "Motivo": 50,
          "Fuente": 32, "Archivo": 30, "Tipo_borrador": 28, "Variedad_fuente": 30, "Avisos_QA": 60, "Categoría": 45}
AJUSTE = {"Español", "Quechua", "Quechua_borrador", "Quechua_validado", "Notas", "Observaciones_heredadas", "Detalle",
          "Antes", "Después", "Motivo", "Avisos_QA", "Escenario", "Otros_borradores", "Evidencia_P6", "Descripción"}


def escribir_hoja(wb, nombre: str, df: pd.DataFrame, estado_col: str | None = "Estado", dominio_col: str | None = "Dominio",
                  anchos: dict | None = None, visible: bool = True):
    ws = wb.create_sheet(nombre[:31])
    cols = list(df.columns)
    ws.append(cols)
    for fila in df.itertuples(index=False):
        ws.append([("" if (v is None or (isinstance(v, float) and v != v)) else v) for v in fila])
    for j, c in enumerate(cols, 1):
        celda = ws.cell(row=1, column=j)
        celda.font = Font(bold=True, color="1F2937")
        celda.fill = ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        ancho = (anchos or {}).get(c) or ANCHOS.get(c) or min(max(10, len(str(c)) + 2), 30)
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"
    idx_dom = cols.index(dominio_col) + 1 if dominio_col and dominio_col in cols else None
    for i in range(2, len(df) + 2):
        relleno = COLOR_DOMINIO.get(ws.cell(row=i, column=idx_dom).value) if idx_dom else None
        for j, c in enumerate(cols, 1):
            celda = ws.cell(row=i, column=j)
            celda.alignment = Alignment(vertical="top", wrap_text=c in AJUSTE)
            celda.border = BORDE
            if relleno is not None:
                celda.fill = relleno
    if estado_col and estado_col in cols and len(df):
        dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"', allow_blank=True)
        letra = get_column_letter(cols.index(estado_col) + 1)
        dv.add(f"{letra}2:{letra}{len(df) + 1}")
        ws.add_data_validation(dv)
    if not visible:
        ws.sheet_state = "hidden"
    configurar_impresion(ws)
    return ws


def configurar_impresion(ws, repetir_encabezado: bool = True):
    """A4 horizontal, ajustado al ancho de la página y con la fila de encabezado repetida al imprimir."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if repetir_encabezado:
        ws.print_title_rows = "1:1"


def escribir_clave_valor(wb, nombre: str, filas: list[tuple], anchos=(40, 90)):
    """Hoja de texto (LEEME / Resumen): filas (clave, valor[, nota]); clave en negrita."""
    ws = wb.create_sheet(nombre[:31])
    for f in filas:
        ws.append(list(f))
    for i, f in enumerate(filas, 1):
        if len(f) == 1 and f[0]:
            ws.cell(row=i, column=1).font = Font(bold=True, size=12, color="1F2937")
        elif f and f[0]:
            ws.cell(row=i, column=1).font = Font(bold=True)
        for j in range(1, 4):
            ws.cell(row=i, column=j).alignment = Alignment(vertical="top", wrap_text=True)
    for j, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = a
    configurar_impresion(ws, repetir_encabezado=False)
    return ws
