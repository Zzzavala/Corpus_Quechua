"""Hoja de grabación de un bloque (§8.6): se genera cuando vuelve el traductor (usa Quechua_validado).

Por sublote: ID (grande), Español, Quechua validado, casillas ☐ ES ☐ QU y espacio para notas.
Uso:
    python scripts/07_hoja_grabacion.py B01                 → 03_bloques/B01/B01_grabacion.docx
    python scripts/07_hoja_grabacion.py B01 --maestro RUTA --salida RUTA [--permitir-faltantes]
"""
from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from comun import CFG, RAIZ, asegurar_dir, guardar_seguro, ruta
import docs_word as W
import planilla as P


def ruta_arg(valor, defecto: Path) -> Path:
    """Rutas relativas: desde la raíz del proyecto (igual que 06 y 08)."""
    if not valor:
        return defecto
    p = Path(valor)
    return p if p.is_absolute() else RAIZ / p


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("bloque")
    ap.add_argument("--maestro")
    ap.add_argument("--salida")
    ap.add_argument("--permitir-faltantes", action="store_true", help="genera igual, marcando las filas aún sin validar")
    a = ap.parse_args()
    b = a.bloque.upper()
    p_maestro = ruta_arg(a.maestro, RAIZ / CFG["rutas"]["maestro"])
    p_salida = ruta_arg(a.salida, ruta("bloques") / b / f"{b}_grabacion.docx")
    filas = [r for r in P.leer(P.abrir(p_maestro)["Corpus"]) if r["Bloque"] == b]
    if not filas:
        print(f"No hay filas del bloque {b} en el maestro.")
        return 2
    filas.sort(key=lambda r: int(r["N_bloque"]))
    faltan = [r["ID"] for r in filas if not r["Quechua_validado"]]
    if faltan and not a.permitir_faltantes:
        print(f"{len(faltan)} filas de {b} aún no tienen quechua validado (p. ej. {', '.join(faltan[:5])}). "
              "Importa primero el Word del traductor (06_importar_traductor.py) o usa --permitir-faltantes.")
        return 3

    doc = W.documento_horizontal(margen_cm=1.2)
    W.pie_con_pagina(doc, f"Grabación {b}")
    doc.add_heading(f"Hoja de grabación — Bloque {b}", level=1)
    au = CFG["audio"]
    fmt = au["formato"]
    canales = "mono" if au["canales"] == 1 else f"{au['canales']} canales"
    for t in ["Lea primero el ESPAÑOL y después el QUECHUA, con una pausa de 2 segundos entre ambos.",
              f"Archivos: {{ID}}_es.{fmt} y {{ID}}_qu.{fmt} ({fmt.upper()}, {au['frecuencia_hz'] / 1000:g} kHz, {canales}, "
              f"{au['bits']} bits). Un sublote por vez.",
              "Si el hablante dice la frase distinto a lo escrito, no se corrige: anote en «Notas» lo que dijo (regla D8).",
              "Marque ☐ ES y ☐ QU cuando cada audio quede bien grabado."]:
        doc.add_paragraph(t, style="List Bullet")
    por_sub = defaultdict(list)
    for r in filas:
        por_sub[r["Sublote"]].append(r)
    anchos = [1.2, 3.2, 8.4, 8.4, 2.4, 3.7]
    for k, s in enumerate(sorted(por_sub)):
        doc.add_page_break()
        fs = por_sub[s]
        doc.add_heading(f"Sublote {k + 1} (filas {fs[0]['N_bloque']}–{fs[-1]['N_bloque']})", level=2)
        tb = W.tabla(doc, ["N°", "ID", "Español", "Quechua", "Grabado", "Notas"], anchos)
        for r in fs:
            fila = tb.add_row()
            W.no_partir_fila(fila)
            c = fila.cells
            W.texto_celda(c[0], r["N_bloque"])
            W.texto_celda(c[1], r["ID"], negrita=True, tamano=14)
            W.texto_celda(c[2], r["Español"], tamano=12)
            qu = r["Quechua_validado"] or "(FALTA VALIDAR)"
            W.texto_celda(c[3], qu, tamano=12, negrita=bool(r["Quechua_validado"]), gris=not r["Quechua_validado"])
            W.texto_celda(c[4], "☐ ES\n☐ QU", tamano=12, alinear=WD_ALIGN_PARAGRAPH.LEFT)
            W.texto_celda(c[5], "")
        W.anchos_columnas(tb, anchos)
    asegurar_dir(p_salida.parent)
    guardar_seguro(p_salida, lambda tmp: doc.save(tmp))
    print(f"OK → {p_salida} ({len(filas)} filas; sin validar: {len(faltan)})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
