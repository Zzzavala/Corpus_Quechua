"""Prueba de los scripts 06, 07 y 08 con datos simulados (§8.5, §13), sin tocar los archivos reales.

Trabaja sobre COPIAS en 02_maestro/checkpoints/prueba_06_08/:
- Copia B01_traductor.docx y simula lo que haría el traductor: deja filas igual, edita un poco, reescribe, llena las
  oraciones de control y las filas sin borrador, deja una vacía, agrega comentarios, cambia un español, borra una fila,
  duplica otra y agrega un ID que no existe.
- Corre 06 (importar) contra copias del maestro y del Excel de trabajo, y verifica POR ID lo que quedó escrito.
- Corre 07 (hoja de grabación) sin y con --permitir-faltantes, y abre el .docx con LibreOffice.
- Crea WAV de prueba (bien, 44.1 kHz estéreo, muy corto, faltante, no-WAV, en subcarpeta) y corre 08 (manifest).
Resultado: 04_reportes/PRUEBA_06_07_08.md (PASA/FALLA por cada chequeo). Código de salida 1 si algo falla.
"""
from __future__ import annotations

import copy
import csv
import math
import shutil
import struct
import subprocess
import sys
import wave
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from comun import CFG, RAIZ, asegurar_dir, escribir_texto, ruta, sha256
import planilla as P
from qa import aviso_word

PY = sys.executable
TMP = RAIZ / "02_maestro" / "checkpoints" / "prueba_06_08"
B = "B01"
REAL_MAESTRO = RAIZ / CFG["rutas"]["maestro"]
REAL_TRABAJO = ruta("bloques") / B / f"{B}_trabajo.xlsx"
REAL_WORD = ruta("bloques") / B / f"{B}_traductor.docx"
chequeos: list[tuple[str, bool, str]] = []


def chequear(nombre, ok, detalle=""):
    chequeos.append((nombre, bool(ok), detalle))
    print(f"[{'PASA ' if ok else 'FALLA'}] {nombre} {detalle}")


def correr(args):
    r = subprocess.run([PY, *args], cwd=RAIZ, capture_output=True, text=True, encoding="utf-8")
    return r.returncode, (r.stdout or "") + (r.stderr or "")


def wav(p: Path, segundos=1.5, fr=16000, canales=1, bits=16):
    asegurar_dir(p.parent)
    n = int(segundos * fr)
    with wave.open(str(p), "wb") as w:
        w.setnchannels(canales)
        w.setsampwidth(bits // 8)
        w.setframerate(fr)
        cuadros = b"".join(struct.pack("<" + "h" * canales, *([int(3000 * math.sin(2 * math.pi * 440 * i / fr))] * canales)) for i in range(n))
        w.writeframes(cuadros)


def main() -> int:
    if TMP.exists():
        shutil.rmtree(TMP)
    asegurar_dir(TMP)
    huellas_antes = {p: sha256(p) for p in (REAL_MAESTRO, REAL_TRABAJO, REAL_WORD)}
    t_maestro, t_trabajo = TMP / "corpus_maestro.xlsx", TMP / f"{B}_trabajo.xlsx"
    t_word = TMP / f"{B}_traductor_devuelto.docx"
    shutil.copy(REAL_MAESTRO, t_maestro)
    shutil.copy(REAL_TRABAJO, t_trabajo)
    shutil.copy(REAL_WORD, t_word)

    wb = P.abrir(t_maestro)
    corpus = {r["ID"]: r for r in P.leer(wb["Corpus"])}
    avisos = {}
    for r in P.leer(wb["QA"]):
        avisos.setdefault(r["ID"], []).append((r["Tipo"], r["Severidad"], r["Detalle"]))

    # ------------------------------------------------------------ 06 --marcar-enviado (al mandar el Word)
    code, out = correr(["scripts/06_importar_traductor.py", B, "--marcar-enviado", "--maestro", str(t_maestro), "--trabajo", str(t_trabajo)])
    wb1 = P.abrir(t_maestro)
    est1 = Counter(r["Estado"] for r in P.leer(wb1["Corpus"]) if r["Bloque"] == B)
    cam1 = [c for c in P.leer(wb1["Cambios"]) if "--marcar-enviado" in c["Regla"]]
    chequear("06 --marcar-enviado deja las 200 filas en Enviado y lo registra en Cambios",
             code == 0 and est1 == Counter({"Enviado": 200}) and len(cam1) == 200, f"({dict(est1)}; {len(cam1)} en Cambios)")

    # ------------------------------------------------------------ simular al traductor
    doc = Document(t_word)
    filas_doc = []
    for t in doc.tables:
        for row in t.rows:
            if len(row.cells) >= 6 and row.cells[1].text.strip()[:4] in ("GEN-", "CLI-"):
                filas_doc.append((t, row))
    ident = lambda r: r.cells[1].text.strip()
    esperado = {}            # ID → quechua que debería quedar validado
    con_borr = [(t, r) for t, r in filas_doc if corpus[ident(r)]["Quechua_borrador"] and corpus[ident(r)]["Control"] != "Sí"]
    control = [(t, r) for t, r in filas_doc if corpus[ident(r)]["Control"] == "Sí"]
    sin_borr = [(t, r) for t, r in filas_doc if not corpus[ident(r)]["Quechua_borrador"]]
    por_sub = defaultdict(list)
    for t, r in con_borr:
        por_sub[corpus[ident(r)]["Sublote"]].append((t, r))
    subs = sorted(por_sub)
    s1 = por_sub[subs[0]]
    plan = {"sin cambios": s1[:8], "editado": s1[8:16], "reescrito": s1[16:19]}
    for s in subs[1:3]:                      # S2 y S3: una edición pequeña cada uno; S4 vuelve sin tocar
        plan["editado"].append(por_sub[s][0])
    sin_tocar = [ident(r) for _, r in por_sub[subs[3]]]
    for _, r in plan["sin cambios"]:
        esperado[ident(r)] = r.cells[4].text
    for _, r in plan["editado"]:
        q = r.cells[4].text
        nuevo = q.rstrip(".?!¡¿") + "lla" + q[len(q.rstrip(".?!¡¿")):]
        r.cells[4].text = nuevo
        esperado[ident(r)] = nuevo
    for k, (_, r) in enumerate(plan["reescrito"]):
        nuevo = f"Huk rimaymi kayqa, musuqmanta qillqasqa {k + 1}."
        r.cells[4].text = nuevo
        esperado[ident(r)] = nuevo
    for k, (_, r) in enumerate(control):
        borr = corpus[ident(r)]["Quechua_borrador"]
        nuevo = borr if k % 2 == 0 else f"Kikinpa qillqasqan rimay {k}."
        r.cells[4].text = nuevo
        esperado[ident(r)] = nuevo
    vacia_id = ident(sin_borr[0][1])
    for k, (_, r) in enumerate(sin_borr[1:], 1):
        nuevo = f"Rimay {k}."
        r.cells[4].text = nuevo
        esperado[ident(r)] = nuevo
    for s in subs[:3]:
        for _, r in por_sub[s]:
            esperado.setdefault(ident(r), r.cells[4].text)
    # comentario después de un aviso y comentario suelto
    con_aviso = next(((t, r) for t, r in con_borr if r.cells[5].text.startswith("Aviso:") and ident(r) not in sin_tocar), None)
    comentario = "Comentario: está bien así, es un préstamo común."
    if con_aviso:
        con_aviso[1].cells[5].text = con_aviso[1].cells[5].text + " " + comentario
    sin_aviso = next((r for t, r in con_borr if not r.cells[5].text and ident(r) not in sin_tocar
                      and ident(r) != ident(plan["sin cambios"][0][1])), None)
    sin_aviso.cells[5].text = "Mejor decir con -taq."
    # el traductor cambia un español (y su quechua): esa fila NO se valida
    cambia_es = plan["sin cambios"][0][1]
    cambia_id = ident(cambia_es)
    cambia_es.cells[3].text = cambia_es.cells[3].text.rstrip(".?!") + " también."
    cambia_es.cells[4].text = cambia_es.cells[4].text + " ari"
    esperado.pop(cambia_id)
    # borrar una fila (que no sea control), duplicar otra, agregar un ID inexistente (en S2)
    s2 = por_sub[subs[1]]
    borrada_t, borrada_r = s2[3]
    borrada_id = ident(borrada_r)
    esperado.pop(borrada_id, None)
    borrada_t._tbl.remove(borrada_r._tr)
    dup_t, dup_r = s2[4]
    dup_id = ident(dup_r)
    dup_r._tr.addnext(copy.deepcopy(dup_r._tr))
    extra = copy.deepcopy(dup_r._tr)
    dup_t._tbl.append(extra)
    filas_t = dup_t.rows
    filas_t[len(filas_t) - 1].cells[1].text = "GEN-9999"
    doc.save(t_word)

    # ------------------------------------------------------------ 06 rechaza un Word con Control de cambios (H2)
    t_cc = TMP / f"{B}_con_control_de_cambios.docx"
    dcc = Document(t_word)
    celda = next(r.cells[4] for t in dcc.tables for r in t.rows if r.cells[1].text.strip()[:4] in ("GEN-", "CLI-"))
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "901")
    ins.set(qn("w:author"), "Traductor")
    ins.set(qn("w:date"), "2026-10-20T10:00:00Z")
    run = OxmlElement("w:r")
    wt = OxmlElement("w:t")
    wt.text = "taq"
    run.append(wt)
    ins.append(run)
    celda.paragraphs[0]._p.append(ins)
    dcc.save(t_cc)
    h_m = sha256(t_maestro)
    code, out = correr(["scripts/06_importar_traductor.py", B, "--docx", str(t_cc), "--maestro", str(t_maestro),
                        "--trabajo", str(t_trabajo), "--reporte", str(TMP / "cc.md")])
    chequear("06 no importa un Word con Control de cambios sin aceptar (y no toca el maestro)",
             code == 4 and sha256(t_maestro) == h_m, out.strip()[:120])

    # ------------------------------------------------------------ 06
    rep = TMP / "importacion_B01.md"
    rel = lambda p: str(p.relative_to(RAIZ))             # rutas relativas a la raíz (H36)
    code, out = correr(["scripts/06_importar_traductor.py", B, "--docx", rel(t_word), "--maestro", rel(t_maestro),
                        "--trabajo", rel(t_trabajo), "--reporte", rel(rep)])
    chequear("06 corre sin error (rutas relativas a la raíz)", code == 0, out.strip().splitlines()[-1] if out.strip() else "")
    wb2 = P.abrir(t_maestro)
    corpus2 = {r["ID"]: r for r in P.leer(wb2["Corpus"])}
    malos = [i for i, q in esperado.items() if corpus2[i]["Quechua_validado"] != q or corpus2[i]["Estado"] != "Validado"]
    chequear("06 escribe el quechua final POR ID en el maestro", not malos and len(esperado) > 120, f"({len(esperado)} filas; mal: {malos[:5]})")
    chequear("06 no valida la fila borrada ni la vacía",
             corpus2[borrada_id]["Quechua_validado"] == "" and corpus2[vacia_id]["Quechua_validado"] == "", f"({borrada_id}, {vacia_id})")
    chequear("06: español cambiado por el traductor → Por_revisar, sin validar (H1)",
             corpus2[cambia_id]["Estado"] == "Por_revisar" and corpus2[cambia_id]["Quechua_validado"] == "", f"({cambia_id})")
    pend = [i for i in sin_tocar if corpus2[i]["Estado"] == "Validado"]
    chequear(f"06: sublote devuelto sin ningún cambio ({subs[3]}) no se valida solo (H7)", not pend and len(sin_tocar) >= 10,
             f"({len(sin_tocar)} filas; validadas por error: {len(pend)})")
    cam2 = [c for c in P.leer(wb2["Cambios"]) if "(06)" in c["Regla"]]
    chequear("06 registra en Cambios lo que importa (H3)",
             sum(1 for c in cam2 if c["Campo"] == "Quechua_validado") == len(esperado) and
             any(c["Campo"] == "Español (propuesta del traductor)" and c["ID"] == cambia_id for c in cam2),
             f"({len(cam2)} registros)")
    # reimportar confirmando el sublote: ahora sí se validan, y nada se registra dos veces
    code, out = correr(["scripts/06_importar_traductor.py", B, "--docx", rel(t_word), "--maestro", rel(t_maestro),
                        "--trabajo", rel(t_trabajo), "--reporte", rel(rep), "--confirmar-sublotes", subs[3]])
    wb3 = P.abrir(t_maestro)
    corpus3 = {r["ID"]: r for r in P.leer(wb3["Corpus"])}
    cam3 = [c for c in P.leer(wb3["Cambios"]) if "(06)" in c["Regla"]]
    dobles = Counter((c["ID"], c["Campo"], c["Después"]) for c in cam3 if c["Campo"] != "Estado")
    chequear("06 --confirmar-sublotes valida el sublote confirmado; reimportar no duplica registros",
             code == 0 and all(corpus3[i]["Estado"] == "Validado" for i in sin_tocar) and max(dobles.values()) == 1,
             f"(registros 06: {len(cam2)} → {len(cam3)})")
    for i in sin_tocar:
        esperado[i] = corpus3[i]["Quechua_validado"]
    from importlib.util import module_from_spec, spec_from_file_location
    spec = spec_from_file_location("imp06", RAIZ / "scripts" / "06_importar_traductor.py")
    imp06 = module_from_spec(spec)
    spec.loader.exec_module(imp06)
    txt0 = imp06.texto_anclaje(0.2, 0.0, 5, 50)
    chequear("06: control con distancia 0 no se confunde con «sin control» (H6)", "distancia 0" in txt0, f"«{txt0[:80]}…»")
    tm = [r for r in P.leer(wb2["Traductor_metricas"]) if r["Bloque"] == B]
    chequear("06 llena Traductor_metricas", len(tm) == 1, "")
    if tm:
        m = tm[0]
        chequear("06 reporta faltantes / sobrantes / duplicados",
                 borrada_id in m["Faltantes"] and "GEN-9999" in m["Sobrantes"] and dup_id in m["Duplicados"],
                 f"(faltantes={m['Faltantes']}; sobrantes={m['Sobrantes']}; duplicados={m['Duplicados']})")
        chequear("06 calcula métricas de edición", float(m["Sin_cambios_%"]) > 0 and float(m["Editado_%"]) > 0 and float(m["Reescrito_%"]) > 0,
                 f"(sin cambios {m['Sin_cambios_%']} %, editado {m['Editado_%']} %, reescrito {m['Reescrito_%']} %, d media {m['Distancia_media_borrador']})")
        chequear("06 analiza las oraciones de control (D7)", int(m["Control_n"]) == len(control) and m["Anclaje"] != "",
                 f"(control {m['Control_n']}; {m['Anclaje']})")
    wbt = P.abrir(t_trabajo)
    det = {r["ID"]: r for r in P.leer(wbt["Detalle"])}
    blo = {r["ID"]: r for r in P.leer(wbt["Bloque"])}
    if con_aviso:
        cid = con_aviso[1].cells[1].text.strip()
        chequear("06 separa el comentario del «Aviso:» prellenado", det[cid]["Comentario_traductor"] == comentario,
                 f"({cid}: «{det[cid]['Comentario_traductor']}»)")
    sid = sin_aviso.cells[1].text.strip()
    chequear("06 guarda el comentario sin aviso", det[sid]["Comentario_traductor"] == "Mejor decir con -taq.", f"({sid})")
    malos_t = [i for i, q in esperado.items() if blo[i]["Quechua"] != q or blo[i]["Estado"] != "Validado"]
    chequear("06 actualiza el Excel de trabajo por ID", not malos_t, f"(mal: {malos_t[:5]})")
    txt = rep.read_text(encoding="utf-8")
    chequear("06 avisa que el traductor cambió un español", cambia_es.cells[1].text.strip() in txt.split("cambió el español en:")[1].split("\n")[0], "")
    huellas_despues = {p: sha256(p) for p in huellas_antes}
    chequear("Los archivos reales no se tocaron", huellas_antes == huellas_despues, "")

    # ------------------------------------------------------------ 07
    t_grab = TMP / f"{B}_grabacion.docx"
    code, out = correr(["scripts/07_hoja_grabacion.py", B, "--maestro", rel(t_maestro), "--salida", rel(t_grab)])
    chequear("07 se niega si faltan filas validadas", code == 3, "")
    code, out = correr(["scripts/07_hoja_grabacion.py", B, "--maestro", rel(t_maestro), "--salida", rel(t_grab), "--permitir-faltantes"])
    chequear("07 genera la hoja de grabación (rutas relativas a la raíz)", code == 0 and t_grab.exists(), out.strip()[-120:])
    if t_grab.exists():
        dg = Document(t_grab)
        n_filas = sum(len(t.rows) - 1 for t in dg.tables)
        chequear("07: 200 filas en 4 tablas (una por sublote)", n_filas == 200 and len(dg.tables) == 4, f"({n_filas} filas, {len(dg.tables)} tablas)")
        code, out = correr(["scripts/render_revision.py", str(t_grab), "--paginas", "2", "--salida", str(TMP / "render")])
        chequear("07: el .docx abre en LibreOffice", code == 0, out.strip()[-160:])

    # ------------------------------------------------------------ 08
    audio = TMP / "audio"
    ids = [i for i in esperado][:8]
    wav(audio / f"{ids[0]}_es.wav"); wav(audio / f"{ids[0]}_qu.wav")
    wav(audio / f"{ids[1]}_es.wav"); wav(audio / f"{ids[1]}_qu.wav", fr=44100, canales=2)
    wav(audio / f"{ids[2]}_es.wav", segundos=0.1); wav(audio / f"{ids[2]}_qu.wav")
    wav(audio / f"{ids[3]}_es.wav")
    wav(audio / "sesion2" / f"{ids[4]}_es.wav"); wav(audio / "sesion2" / f"{ids[4]}_qu.wav")
    (audio / f"{ids[5]}_es.wav").write_bytes(b"esto no es un wav")
    wav(audio / f"{ids[5]}_qu.wav")
    wav(audio / f"{ids[6]}_ES.WAV"); wav(audio / f"{ids[6]}_qu.wav")            # nombre en mayúsculas
    wav(audio / f"{ids[7]}_es.wav"); wav(audio / f"{ids[7]}_qu.wav"); wav(audio / "sesion3" / f"{ids[7]}_qu.wav")   # repetido
    man = TMP / "manifest_B01.csv"
    code, out = correr(["scripts/08_manifest.py", B, "--maestro", rel(t_maestro), "--audio-dir", rel(audio), "--salida", rel(man)])
    chequear("08 corre y escribe el manifest (rutas relativas a la raíz)", code == 0 and man.exists(), out.strip()[-140:])
    filas = {r["id"]: r for r in csv.DictReader(open(man, encoding="utf-8"))}
    chequear("08: 200 filas con las columnas pedidas", len(filas) == 200 and list(next(iter(filas.values())).keys()) ==
             ["id", "bloque", "sublote", "split", "dominio", "rol", "es", "qu", "audio_es", "audio_qu", "dur_es", "dur_qu", "qa_audio"], "")
    chequear("08: audio correcto → OK", filas[ids[0]]["qa_audio"] == "OK" and filas[ids[4]]["qa_audio"] == "OK", f"({filas[ids[0]]['qa_audio']}; {filas[ids[4]]['qa_audio']})")
    chequear("08: detecta 44.1 kHz y estéreo", "44100 Hz" in filas[ids[1]]["qa_audio"] and "2 canales" in filas[ids[1]]["qa_audio"], filas[ids[1]]["qa_audio"])
    chequear("08: detecta audio muy corto", "muy corto" in filas[ids[2]]["qa_audio"], filas[ids[2]]["qa_audio"])
    chequear("08: detecta audio faltante", "FALTA_QU" in filas[ids[3]]["qa_audio"], filas[ids[3]]["qa_audio"])
    chequear("08: detecta archivo que no es WAV", "no es WAV" in filas[ids[5]]["qa_audio"], filas[ids[5]]["qa_audio"])
    chequear("08: marca filas sin quechua validado", "SIN_VALIDAR" in filas[vacia_id]["qa_audio"], "")
    chequear("08: encuentra el audio aunque el nombre esté en mayúsculas", filas[ids[6]]["qa_audio"] == "OK", filas[ids[6]]["qa_audio"])
    chequear("08: avisa audios con el mismo nombre en dos carpetas (H33)", "AUDIO_REPETIDO_QU" in filas[ids[7]]["qa_audio"],
             filas[ids[7]]["qa_audio"][:120])

    ok = all(c[1] for c in chequeos)
    md = ["# Prueba de los scripts 06, 07 y 08 con datos simulados", "",
          f"Generado por `scripts/test_ida_vuelta.py`. Se trabajó sobre copias en `02_maestro/checkpoints/prueba_06_08/`; "
          "los archivos reales no se modificaron (se comprueba con sha256).", "",
          "**Qué se simuló:** marcar el bloque como Enviado; en la copia del Word de B01, filas sin cambios, 10 con una edición "
          "pequeña (8 en S1 y 1 en S2 y S3), 3 reescritas, el sublote S4 devuelto sin ningún cambio, "
          f"las {len(control)} oraciones de control llenadas (la mitad igual al borrador oculto, la otra mitad distinta), "
          "las filas sin borrador llenadas, 1 fila dejada vacía, un comentario después de un «Aviso:», un comentario suelto, "
          "un español cambiado por el traductor, 1 fila borrada, 1 fila duplicada y 1 ID inexistente (GEN-9999); además un Word "
          "con Control de cambios sin aceptar y una reimportación confirmando S4. "
          "Para 08 se crearon WAV de prueba: correctos, a 44.1 kHz estéreo, muy corto, faltante, un archivo que no es WAV, "
          "uno en subcarpeta, uno con el nombre en mayúsculas y uno repetido en dos carpetas.", "",
          f"**Resultado: {'TODO PASA' if ok else 'HAY FALLAS'}** ({sum(c[1] for c in chequeos)} de {len(chequeos)}).", "",
          "| Chequeo | Resultado | Detalle |", "|---|---|---|"]
    md += [f"| {n} | {'PASA' if o else 'FALLA'} | {d.replace('|', '/')} |" for n, o, d in chequeos]
    escribir_texto(ruta("reportes") / "PRUEBA_06_07_08.md", "\n".join(md))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
