"""Ayudas para los Word (python-docx): A4 horizontal, letra legible, tablas con encabezado repetido."""
from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GRIS = RGBColor(0x6B, 0x72, 0x80)
FONDO_ENCABEZADO = "DCE8F7"


def documento_horizontal(margen_cm: float = 1.2, letra: str = "Calibri", tamano: int = 11) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(margen_cm))
    estilo = doc.styles["Normal"]
    estilo.font.name = letra
    estilo.font.size = Pt(tamano)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), letra)
    return doc


def documento_vertical(margen_cm: float = 2.0, letra: str = "Calibri", tamano: int = 11) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(margen_cm))
    estilo = doc.styles["Normal"]
    estilo.font.name = letra
    estilo.font.size = Pt(tamano)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), letra)
    return doc


def pie_con_pagina(doc: Document, texto: str):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(texto + " — página ")
    r.font.size = Pt(9)
    r.font.color.rgb = GRIS
    for tipo, instr in (("begin", None), (None, "PAGE"), ("end", None)):
        run = p.add_run()
        run.font.size = Pt(9)
        run.font.color.rgb = GRIS
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            run._r.append(it)


def sombrear(celda, color: str):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def repetir_encabezado(fila):
    trPr = fila._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


def no_partir_fila(fila):
    trPr = fila._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    cs.set(qn("w:val"), "true")
    trPr.append(cs)


def anchos_columnas(tabla, anchos_cm: list[float]):
    tabla.autofit = False
    tbl = tabla._tbl
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), anchos_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for fila in tabla.rows:
        for celda, w in zip(fila.cells, anchos_cm):
            celda.width = Cm(w)


def texto_celda(celda, texto: str, negrita=False, cursiva=False, gris=False, tamano: float | None = None, alinear=None):
    celda.text = ""
    p = celda.paragraphs[0]
    if alinear:
        p.alignment = alinear
    r = p.add_run(texto)
    r.bold = negrita
    r.italic = cursiva
    if gris:
        r.font.color.rgb = GRIS
    if tamano:
        r.font.size = Pt(tamano)
    return r


def tabla(doc: Document, encabezados: list[str], anchos_cm: list[float], estilo="Table Grid"):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = estilo
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(t.rows[0].cells, encabezados):
        texto_celda(c, h, negrita=True)
        sombrear(c, FONDO_ENCABEZADO)
    repetir_encabezado(t.rows[0])
    anchos_columnas(t, anchos_cm)
    return t
