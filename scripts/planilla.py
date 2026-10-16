"""Leer y actualizar celdas de un .xlsx ya formateado (sin perder el formato), buscando columnas por nombre."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from comun import guardar_seguro


def abrir(path: Path):
    try:
        return load_workbook(path)
    except PermissionError as e:
        raise PermissionError(f"No se puede abrir {path}: ¿está abierto en Excel/LibreOffice? Ciérralo y vuelve a correr.") from e


def columnas(ws) -> dict[str, int]:
    return {str(c.value): c.column for c in ws[1] if c.value is not None}


def filas_por_id(ws, col_id: str = "ID") -> dict[str, int]:
    cid = columnas(ws)[col_id]
    out = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=cid).value
        if v:
            out[str(v)] = r
    return out


def leer(ws) -> list[dict]:
    cols = [c.value for c in ws[1]]
    datos = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if all(v in (None, "") for v in fila):
            continue
        datos.append({str(k): ("" if v is None else str(v)) for k, v in zip(cols, fila) if k is not None})
    return datos


def poner(ws, fila: int, col: str, valor):
    ws.cell(row=fila, column=columnas(ws)[col]).value = valor


def guardar(wb, path: Path):
    guardar_seguro(path, lambda tmp: wb.save(tmp))
